r"""Gera o GUIA DAS SKILLS em Word (claude-kit\GUIA_DAS_SKILLS.docx).

Ettore, 06/10: "um arquivo em Word para saber o que cada skill faz e quando usá-las, se as uso
ativamente, com o comando '/', ou se as uso dentro do próprio projeto, já programado pelo CLAUDE.md
ou outros documentos acessados apenas quando necessários".

O TEXTO vive AQUI, versionado; o .docx é GERADO. Mudou uma skill, um plugin ou um projeto? Rode de
novo (e, se aparecer skill em "Sem descrição curada", acrescente-a ao dicionário CURADO, abaixo):

    python C:\CLAUDE-PROJETOS\claude-kit\scripts\gerar_guia_skills.py
    python C:\CLAUDE-PROJETOS\claude-kit\scripts\gerar_guia_skills.py --saida C:\algum\outro.docx

O gerador LÊ o que está instalado de verdade, para o guia nunca mentir:
  - as skills do kit, por grupo (claude-kit\plugins\<grupo>\skills\*\SKILL.md, pelo
    .claude-plugin\marketplace.json), e a conferência do instalar_kit.py --verificar;
  - os plugins de fora ligados (enabledPlugins do settings.json) e as skills e comandos deles, no cache;
  - que grupo cada projeto liga (enabledPlugins do usuário, de cada <projeto>\.claude\settings*.json e
    da raiz), para a tabela projeto × grupo;
  - a operação (COMO_OPERAR.md), que vira a 1ª parte do guia;
  - as skills sincronizadas do claude.ai (~\.claude\skills\synced\*\manifest.json);
  - as skills e os comandos dos projetos (<projeto>\.claude\skills e \.claude\commands): só o nome
    e a descrição. NUNCA abre as pastas de dados (PROIBIDAS, abaixo);
  - as do próprio Claude Code, que não dá para ler do disco: lista fixa aqui (NATIVAS).

Do disco vêm: o nome, a descrição e a marca "só /" (disable-model-invocation). Escritos à mão, no
CURADO: o grupo, o "para quê", o "quando usar" e o exemplo de pedido. Skill instalada sem texto à
mão entra em "Sem descrição curada"; texto à mão de skill que saiu do PC simplesmente não aparece.

Referência a outra skill, dentro de um texto: [[nome]] (vira `/nome` se só se chama por /, e `nome`
se o Claude a chama sozinho; avisa se a skill citada não está instalada).

Precisa de python-docx. Se o Word estiver com o .docx aberto, o gerador não força: grava com outro
nome e avisa.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

# ----------------------------------------------------------------------
# onde as coisas estão (tudo sai da posição deste arquivo e da pasta do usuário)
# ----------------------------------------------------------------------
AQUI = Path(__file__).resolve()
KIT = AQUI.parents[1]                      # claude-kit
RAIZ_PROJETOS = KIT.parent                 # C:\CLAUDE-PROJETOS
CLAUDE_HOME = Path.home() / '.claude'
DESTINO = KIT / 'GUIA_DAS_SKILLS.docx'
COMANDO = f'python {AQUI}'

# Projetos lidos (pasta -> como o Ettore os chama). Projeto novo com skills: acrescente aqui e
# escreva o texto no CURADO. (Pasta com .claude\skills ou .claude\commands que não esteja aqui é
# achada pelo gerador e aparece em "Sem descrição curada".)
PROJETOS = {
    'desosp-app': 'app',
    'desosp-censo': 'kernel',
    'desosp-hc': 'planilha HC',
}
# Pastas de dados: o gerador nunca lê dentro delas (confere cada caminho antes de abrir).
PROIBIDAS = {'entrada', 'saida', 'historico', 'arquivo', 'privado', 'planilhas', 'trabalho',
             'desosp-app-dados', 'desosp-app-backups'}

# As skills do próprio Claude Code (não dá para ler do disco). Conferidas na lista de skills de uma
# sessão de 06/10/2026: todas aparecem para o Claude chamar, logo todas são "sozinha ou /".
NATIVAS = ['code-review', 'simplify', 'security-review', 'init', 'loop', 'schedule', 'dataviz',
           'claude-api', 'fewer-permission-prompts', 'run']

AVISOS: list[str] = []          # tudo o que o gerador achou fora do esperado (vai ao guia e à tela)


# ----------------------------------------------------------------------
# O DICIONÁRIO CURADO — o texto à mão. A ordem aqui é a ordem em cada tabela.
#   g: grupo · p: para quê (1 linha; '@disco' = a descrição lida do arquivo) · q: quando usar
#   e: exemplo de pedido (começa com / = digitado; senão, é frase para dizer ao Claude)
#   t: marca (aparece sob o nome) · s: 'desligar' ou 'decidir' (skills do claude.ai)
# Chave: o nome da skill; nas de projeto, 'pasta:nome' (ex.: 'desosp-hc:rodada').
# ----------------------------------------------------------------------
GRUPOS = [
    # os grupos do kit (plugins de claude-kit\plugins): a skill do kit cai no grupo do plugin dela
    ('nucleo', 'Grupo nucleo — as nossas, em todo projeto',
     'Como usar o Claude Code (modelo, esforço, janelas, contexto, pesquisa) e que recursos um projeto '
     'usa. Ligado no escopo de usuário: vale em toda janela de todo projeto.'),
    ('planejamento', 'Grupo planejamento — entrevistar, decidir, registrar (terceiros), em todo projeto',
     'A família do grill (Matt Pocock) e o que vem junto: perguntas, glossário e decisões, aula, '
     'retrospectiva, escrever para agentes, o /goal. Ligado no escopo de usuário.'),
    ('engenharia', 'Grupo engenharia — do plano ao código (terceiros), só nos projetos de código',
     'Especificação, tarefas, implementação, testes, diagnóstico e interface. Cada projeto de código liga '
     'este grupo no seu `.claude\\settings.json`, na fase dele; fora dali estas skills não aparecem.'),
    ('kit', 'Grupo kit — as nossas, só no kit e na janela da raiz',
     'Organizar os projetos de um PC, as cores e os ícones das pastas. Só com `/`.'),
    ('sdd', 'Grupo sdd — executar um plano de várias tarefas (terceiros), só nos projetos de código',
     'O subagent-driven-development do Superpowers e as 3 skills de que ele depende: um subagente novo por '
     'tarefa, com revisão. Ligado por projeto de código, junto com a engenharia; a nossa `orquestrar` diz '
     'quando usar.'),
    ('pensar','Pensar e decidir (a família do grill, de Matt Pocock)',
     'Antes de construir, decidir. A família do grill faz perguntas, uma de cada vez, até não '
     'sobrar decisão escondida.'),
    ('codigo', 'Do plano ao código',
     'Da especificação ao código que funciona: tarefas, implementação, testes e revisão. As '
     'marcadas “Claude Code” já vêm de fábrica.'),
    ('interface', 'Interface',
     'Para criar e revisar telas e gráficos.'),
    ('janelas', 'Janelas, contexto e memória',
     'Como organizar o trabalho entre janelas (conversas), o que levar de uma para a outra e como '
     'gastar menos contexto.'),
    ('controles', 'Controles',
     'O que protege você de comandos perigosos e de perguntas de permissão demais.'),
    ('rotinas', 'Rotinas e agendamento',
     'Para o Claude repetir uma tarefa sozinho, de tempos em tempos.'),
    ('documentos', 'Documentos (do claude.ai)',
     'Vêm do claude.ai (sincronizadas) e valem em todos os projetos. A descrição de cada uma entra '
     'em toda conversa: por isso as que você não usa valem ser desligadas, lá no claude.ai.'),
    ('projeto', 'Do projeto',
     'Só existem dentro de um projeto: aparecem quando a janela é aberta naquela pasta. O '
     '`CLAUDE.md` do projeto costuma mandar usá-las.'),
]

CURADO: dict[str, dict] = {}


def _c(chave, g, p, q, e=None, t=None, s=None):
    CURADO[chave] = dict(g=g, p=p, q=q, e=e, t=t, s=s)


# --- Pensar e decidir -------------------------------------------------
_c('grill-me', 'pensar',
   'Entrevista em rodadas, uma pergunta por vez, até não sobrar decisão em aberto.',
   'Ideia nova, plano ou decisão com várias partes, antes de construir qualquer coisa.',
   '/grill-me quero criar uma tela de pendências no app')
_c('grill-with-docs', 'pensar',
   'A mesma entrevista, gravando o glossário (`GLOSSARY.md`) e as decisões (`docs/adr/`) enquanto '
   'conversa.',
   'Assunto com vocabulário próprio, ou decisão que precisa ficar escrita para as próximas '
   'janelas.',
   '/grill-with-docs vamos definir o que é “rodada” neste projeto')
_c('grilling', 'pensar',
   'O motor das duas anteriores: o jeito de entrevistar, com a recomendação em cada pergunta.',
   'Entra sozinha quando você pede “me questione” ou “teste este plano”, sem digitar nada.',
   '“Quero testar este plano: me questione antes de eu decidir.”')
_c('domain-modeling', 'pensar',
   'Cuida do glossário e das decisões registradas; desafia termo ambíguo.',
   'Ao discutir o nome de um conceito ou ao registrar uma decisão. Entra sozinha quando o '
   'assunto é o vocabulário do projeto.',
   '“Anote no glossário o que quer dizer ‘alta prevista’.”')
_c('prototype', 'pensar',
   'Protótipo descartável, para responder uma pergunta que só se responde vendo.',
   'Dúvida de tela (“como ficaria?”) ou de lógica (“esse cálculo faz sentido?”), antes de '
   'decidir.',
   '“Faça duas versões da tela de pendências para eu comparar.”')
_c('to-questionnaire', 'pensar',
   'Transforma em questionário a decisão que depende de outra pessoa.',
   'Quando quem sabe a resposta não é você: você envia o questionário e as respostas voltam.',
   '/to-questionnaire o que preciso perguntar à equipe antes de decidir isto')
_c('wait-what', 'pensar',
   'Reexplica em linguagem simples o que o Claude acabou de dizer e você não entendeu.',
   'Quando a última resposta ficou difícil. Digite só o nome e ele refaz, com mais clareza.',
   '/wait-what')
_c('teach', 'pensar',
   'Aula sobre um conceito, usando o seu próprio projeto como exemplo.',
   'Quando você quer aprender (e não só pedir o código): como funciona um teste, o que é uma API.',
   '/teach o que é um teste automático')

# --- Do plano ao código -----------------------------------------------
_c('to-spec', 'codigo',
   'Transforma a conversa em especificação (o que construir e por quê) e a grava onde o projeto '
   'guarda as tarefas. Sem nova entrevista.',
   'Logo depois do grill, na mesma janela, com o entendimento já confirmado. Em projeto novo, '
   'antes dela vem a [[recursos-do-projeto]].',
   '/to-spec')
_c('to-tickets', 'codigo',
   'Divide a especificação em tarefas pequenas, cada uma dizendo do que depende.',
   'Depois do [[to-spec]]. Se o trabalho for grande demais para uma janela, use o [[wayfinder]].',
   '/to-tickets')
_c('implement', 'codigo',
   'Faz uma tarefa por vez: escreve o teste primeiro, programa, revisa e salva uma versão '
   '(commit).',
   'Com a especificação e as tarefas prontas, para construir de verdade.',
   '/implement a próxima tarefa')
_c('implement-spec', 'codigo',
   'O mesmo, em paralelo: um subagente (ajudante) por tarefa, depois junta tudo.',
   'Só trabalho grande que se divide em tarefas independentes. É caro: gasta bem mais que o '
   '[[implement]].',
   '/implement-spec')
_c('setup-matt-pocock-skills', 'codigo',
   'Configura o projeto para estas skills: onde ficam as tarefas e como se escrevem os rótulos.',
   'Uma vez por projeto, antes de usar [[to-spec]], [[to-tickets]] ou [[implement]] pela primeira '
   'vez.',
   '/setup-matt-pocock-skills')
_c('wayfinder', 'codigo',
   'Mapa de decisões para um trabalho maior que uma janela, resolvido aos poucos, ao longo de '
   'várias sessões.',
   'Quando a especificação ficou grande demais e não cabe numa janela só.',
   '/wayfinder reescrever o fluxo de pendências do app')
_c('tdd', 'codigo',
   'Escreve primeiro o teste que falha, depois o código que o faz passar.',
   'Funcionalidade nova ou correção de erro cujo resultado pode ser conferido por teste. Entra '
   'sozinha quando você fala em “teste primeiro”.',
   '“Corrija esse cálculo de dias com tdd.”')
_c('diagnosing-bugs', 'codigo',
   'Roteiro para bug difícil: reproduzir o erro, levantar hipóteses e eliminar uma a uma.',
   'Algo quebrou, dá erro ou ficou lento e a causa não é óbvia. Entra sozinha quando você diz '
   '“diagnostique”.',
   '“A rodada para no merge e não sei por quê: diagnostique.”')
_c('property-based-testing', 'codigo',
   'Testes que varrem um domínio inteiro de entradas (datas, âncoras), em vez de uns poucos '
   'exemplos.',
   'Código que mexe com datas, textos ou conversões, cuja regra vale para qualquer entrada.',
   '“Escreva testes de propriedade para a função da data âncora.”')
_c('spec-to-code-compliance', 'codigo',
   'Confere o código contra o documento que o especifica: o que bate, o que contradiz, o que '
   'falta e o que o código faz sem estar escrito.',
   'Para saber se o código segue o documento que o descreve (um contrato, uma especificação).',
   '“Confira se o merge segue o contrato da fonte.”')
_c('codebase-design', 'codigo',
   'Vocabulário para desenhar módulos: onde cortar, o que esconder, o que mostrar.',
   'Ao desenhar ou melhorar um módulo, ou para deixar o código mais fácil de testar e de o '
   'Claude se achar nele.',
   '“Onde devo cortar este módulo para testar sem o banco?”')
_c('improve-codebase-architecture', 'codigo',
   'Varre o código, mostra num mapa visual (página HTML) onde a arquitetura pode melhorar e '
   'entrevista você sobre o ponto escolhido.',
   'Quando o código cresceu e ficou difícil de mexer. Vale uma vez por fase, não a cada tarefa.',
   '/improve-codebase-architecture')
_c('code-review', 'codigo',
   'Revisa o que mudou (ou uma PR, o pedido de mesclagem no GitHub) atrás de erros, no nível de '
   'esforço pedido.',
   'Antes de aceitar uma mudança grande, ou ao fechar uma fase. Aceita --fix para já corrigir.',
   '/code-review')
_c('simplify', 'codigo',
   'Revisa o que mudou atrás de código repetido ou complicado e já aplica a simplificação.',
   'Depois de implementar, antes do commit. Não procura erro: para isso, o [[code-review]].',
   '/simplify')
_c('security-review', 'codigo',
   'Revisão de segurança das mudanças pendentes na branch atual.',
   'Antes de subir mudanças que mexem com login, arquivos, dados ou entrada do usuário.',
   '/security-review')
_c('run', 'codigo',
   'Abre e conduz o app do projeto, para você ver a mudança funcionando de verdade.',
   'Para conferir no app real (e não só nos testes) que a mudança funciona.',
   '/run')
_c('claude-api', 'codigo',
   'Referência da API do Claude: modelos, preços, parâmetros, cache de prompt.',
   'Só quando for programar algo que chama a API da Anthropic, ou ao perguntar preço e limite '
   'de modelo. Entra sozinha quando o pedido cita Claude ou Anthropic.',
   '“Quanto custa o modelo X na API?”')

# --- Interface --------------------------------------------------------
_c('frontend-design', 'interface',
   'Direção visual para tela nova: tipografia, cor e composição, sem cara de “IA genérica”.',
   'Ao criar ou redesenhar uma tela, antes de programar o visual. Entra sozinha em pedidos de '
   'tela.',
   '“Desenhe a tela de pendências com direção visual própria.”')
_c('web-design-guidelines', 'interface',
   'Revisa o código da tela contra diretrizes de interface (acessibilidade, foco, formulários). '
   'Baixa a lista atual de regras.',
   'Com a tela pronta, para auditar acessibilidade e boas práticas. Entra sozinha quando você '
   'pede “revise a interface”.',
   '“Revise esta tela contra as diretrizes de interface.”')
_c('dataviz', 'interface',
   'Gráfico e painel acessíveis: cores, rótulos, legenda e tema claro e escuro.',
   'Sempre que for fazer um gráfico, painel ou indicador. Entra sozinha.',
   '“Faça um gráfico de altas por dia que um daltônico consiga ler.”')

# --- Janelas, contexto e memória -------------------------------------
_c('uso-do-claude', 'janelas',
   'O núcleo: a escada de modelo e esforço, os sinais de que o modelo não deu conta, as armadilhas '
   '(o Sonnet 5 por engano, o `/effort` que grava), o contexto, uma janela por fase e o planejamento, com a '
   'rodada de contingências. Aponta as quatro peças abaixo.',
   'Entra sozinha ao abrir uma janela, ao planejar e quando você pergunta por que o uso está alto. O '
   '`CLAUDE.md` pessoal manda usá-la.',
   '“Por que esta janela está gastando tanto?”')
_c('modelo-e-esforco', 'janelas',
   'Qual modelo e qual `/effort` para cada tipo de tarefa, quando subir ou descer, e a 1ª linha do '
   'prompt.',
   'Entra sozinha quando o Claude vai recomendar modelo ou esforço.',
   '“Qual modelo e esforço uso nesta fase?”')
_c('fechar-janela', 'janelas',
   'Fecha a janela: reescreve o `docs\\ESTADO.md`, grava o `docs\\PROXIMO.md` (o prompt da próxima) e a '
   'fila de perguntas, faz commit e push e entrega o resumo com os sinais. Traz os modelos desses arquivos.',
   'No fim de cada fase, ou perto de ~150 mil tokens. Entra sozinha.',
   '“Feche a janela.”')
_c('pesquisa', 'janelas',
   'Antes de pesquisar, procura na biblioteca do kit (`pesquisas\\`); o que falta, pesquisa por '
   'subagente; depois, salva lá.',
   'Entra sozinha antes de pesquisa na internet ou varredura grande.',
   '“Pesquise como o Claude Code trata ganchos em plugins.”')
_c('orquestrar', 'janelas',
   'Delegar e paralelizar com economia: subagentes por papel, com modelo e esforço fixados; SDD; '
   'Workflow (teto 10); várias janelas; a execução largada com a fila de perguntas.',
   'Entra sozinha ao delegar, ao executar um plano ou antes de trabalho longo sem supervisão.',
   '“Execute o plano das 5 tarefas.”')
_c('organizar-projetos', 'kit',
   'Reunir os projetos e o kit numa pasta-mãe, pôr no GitHub o que falta (privado), reaproveitar o '
   'que o PC já tem e criar projeto novo.',
   'PC novo, projetos espalhados, ou projeto novo (a pasta nasce na janela da raiz). Só com `/`.',
   '/organizar-projetos')
_c('cores-das-pastas', 'kit',
   'O sistema de cores das pastas: aplicar, conferir, marcar a pasta provisória, legenda, mapa e '
   'VS Code. O gancho que pinta a pasta nova é do grupo nucleo e roda em todo projeto.',
   'Depois de mudar o mapa, num projeto novo, ou quando o `conferir` acusar algo. Só com `/`.',
   '/cores-das-pastas')
_c('icones', 'kit',
   'Os desenhos dos ícones, a montagem do `.ico` em cada PC e a prévia. Pedido de ícone novo se '
   'anota e se desenha em lote.',
   'Quando surgir um verbo ou marca nova, ou um desenho a trocar. Só com `/`.',
   '/icones')
_c('subagent-driven-development', 'codigo',
   'Executa um plano de várias tarefas: um subagente novo por tarefa, em série, com revisão de '
   'especificação e de qualidade, e uma revisão final.',
   'Plano de 4 ou mais tarefas no mesmo código, com cabeçalhos “Task N”. A `orquestrar` diz modelo e '
   'esforço de cada despacho.',
   '“Execute o plano com o subagent-driven-development.”')
_c('requesting-code-review', 'codigo',
   'Pede uma revisão de código a um subagente revisor, com o pacote do que mudou.',
   'Chamada pelo [[subagent-driven-development]]; também ao fechar uma tarefa importante.',
   '“Peça uma revisão deste ramo.”')
_c('using-git-worktrees', 'codigo',
   'Cria ou confere um worktree (uma cópia de trabalho isolada do repositório) antes de mexer.',
   'Chamada pelo [[subagent-driven-development]]; também quando duas janelas precisam do mesmo projeto.',
   '“Abra um worktree para esta tarefa.”')
_c('finishing-a-development-branch', 'codigo',
   'Fecha o ramo: testes, e a escolha entre mesclar, abrir PR ou manter. Merge e push seguem a Q10.',
   'No fim do [[subagent-driven-development]].',
   '“Termine este ramo.”')
_c('recursos-do-projeto', 'janelas',
   'Depois do grill de um projeto novo: decide que ferramentas entram e quais ficam de fora, e '
   'escreve as regras do projeto.',
   'Projeto novo, ou implementação nova, entre o grill e o plano. O `CLAUDE.md` pessoal já manda '
   'usá-la nesse ponto.',
   '“Que ferramentas este projeto novo deve usar?”',
   t='em reavaliação')
_c('claude-handoff', 'janelas',
   'Passa o trabalho a uma sessão nova, que começa na hora, em segundo plano (`claude --bg`).',
   'No fim de uma fase, para a próxima janela continuar com o contexto limpo. Ainda sem teste '
   'neste PC.',
   '/claude-handoff continuar a próxima fase',
   t='experimental')
_c('chief-of-staff', 'janelas',
   'Uma janela-chefe que só coordena: delega o trabalho a subagentes.',
   'Sessão longa com muitas partes, supervisionada por você.',
   '/chief-of-staff',
   t='experimental')
_c('goal-prompt', 'janelas',
   'Escreve a condição do `/goal` (o objetivo de uma sessão longa), pronta para copiar e colar.',
   'Antes de deixar uma sessão longa trabalhando sozinha.',
   '“Escreva um /goal para auditar as regras do kernel.”')
_c('retro', 'janelas',
   'Retrospectiva da sessão: o que funcionou, o que travou e o que mudar.',
   'No fim de uma fase ou janela (passo 6 do fluxo do grill).',
   '/retro')
_c('claude-md-improver', 'janelas',
   'Audita os `CLAUDE.md`: confere a qualidade contra um modelo, mostra o relatório e faz ajustes '
   'pontuais.',
   'Quando o `CLAUDE.md` de um projeto cresceu demais ou envelheceu (alvo da casa: menos de 200 '
   'linhas). Entra sozinha quando você pede para auditar.',
   '“Audite o CLAUDE.md deste projeto.”')
_c('revise-claude-md', 'janelas',
   'Grava no `CLAUDE.md` o que esta sessão aprendeu sobre o projeto.',
   'No fim de uma janela em que você corrigiu o Claude ou descobriu uma regra nova. O `CLAUDE.md` '
   'pessoal tem dois endereços (hardlink): depois de editar, confira com `fsutil hardlink list`.',
   '/revise-claude-md')
_c('session-report', 'janelas',
   'Gera uma página HTML com o uso de tokens, o cache e os pedidos mais caros das suas sessões.',
   'Quando o uso estiver alto e você quiser saber por quê. O arquivo cai na pasta da janela: '
   'guarde fora do git.',
   '“Gere o relatório de uso dos últimos 7 dias.”')
_c('writing-for-agents', 'janelas',
   'Estilo para escrever skills e `CLAUDE.md`: curtos, diretos, sem encher o contexto.',
   'Ao criar ou editar uma skill, um `CLAUDE.md` ou um `AGENTS.md`. Entra sozinha nesses casos.',
   '“Reescreva esta skill no estilo certo.”')
_c('init', 'janelas',
   'Cria o `CLAUDE.md` de um projeto a partir do que encontra no código.',
   'Ao abrir um projeto que ainda não tem `CLAUDE.md`. Depois, enxugue.',
   '/init')

# --- Controles --------------------------------------------------------
_c('cc-safety-net', 'controles',
   'A skill explica por que um comando foi barrado e ajuda a ajustar as regras. Quem barra é o '
   'gancho do plugin, que roda sozinho.',
   'Quando o Claude avisar que um comando foi bloqueado e você quiser entender (ou achar que foi '
   'engano).',
   '/cc-safety-net por que o último comando foi barrado?')
_c('fewer-permission-prompts', 'controles',
   'Lê suas conversas antigas e libera de vez os comandos de leitura mais comuns, para o Claude '
   'perguntar menos.',
   'Quando as perguntas de permissão estiverem atrapalhando. Ele grava em `.claude/settings.json` '
   'do projeto: confira depois o que entrou.',
   '/fewer-permission-prompts')

# --- Rotinas e agendamento -------------------------------------------
_c('loop', 'rotinas',
   'Repete um pedido ou comando a cada intervalo (por exemplo, a cada 5 minutos).',
   'Para acompanhar algo que demora (um processo, um status) sem ficar perguntando.',
   '/loop 5m confira se os testes terminaram')
_c('schedule', 'rotinas',
   'Cria rotinas: agentes na nuvem que rodam num horário marcado, como um despertador com tarefa.',
   'Tarefa que deve rodar sozinha, depois ou toda semana. Roda na nuvem: nunca com dado de '
   'paciente.',
   '/schedule')

# --- Documentos (do claude.ai) ---------------------------------------
_c('xlsx', 'documentos',
   'Cria, lê e corrige planilhas Excel (`.xlsx`, `.csv`), com fórmulas e formatação.',
   'Qualquer tarefa em que a planilha é o arquivo de entrada ou de saída. Entra sozinha quando '
   'você cita a planilha.',
   '“Abra esta planilha e acrescente uma coluna com o total por hospital.”')
_c('docx', 'documentos',
   'Cria e edita documentos Word (`.docx`): títulos, tabelas, comentários e revisões.',
   'Quando o resultado precisa ser um arquivo Word, ou você entrega um.',
   '“Gere em Word o guia desta rodada.”')
_c('pdf', 'documentos',
   'Lê, junta, divide, preenche formulário e cria PDF; lê também PDF escaneado (OCR).',
   'Qualquer tarefa com arquivo `.pdf`.',
   '“Junte estes três PDFs em um só.”')
_c('pptx', 'documentos',
   'Cria e edita apresentações PowerPoint (`.pptx`).',
   'Quando a entrega é um arquivo `.pptx`.',
   '“Faça uma apresentação de 5 slides com estes resultados.”')
_c('docs', 'documentos',
   'Documentos editáveis no próprio Claude (Claude Docs): você e outras pessoas editam e '
   'comentam; exportam para Word, PDF e Google Docs.',
   'Para qualquer texto que você vai guardar, compartilhar ou comentar (relatório, carta, guia) '
   'e que não precisa ser arquivo Word.',
   '“Escreva a proposta como um doc.”')
_c('skill-creator', 'documentos',
   'Cria e melhora skills, mede o desempenho e afina a descrição para a skill disparar na hora '
   'certa.',
   'Ao criar uma skill sua ou testar se uma skill funciona como deveria.',
   '“Teste se a descrição desta skill dispara na hora certa.”')
_c('stop-slop', 'documentos',
   'Tira o “jeito de IA” do texto: frases batidas, enchimento, tom artificial.',
   'Ao escrever e-mail, mensagem ou texto para outras pessoas.',
   '“Reescreva este e-mail sem cara de IA.”')
_c('task-observer', 'documentos',
   'Observa as tarefas e anota ideias de skill nova. A descrição pede para ser chamada no começo '
   'de toda tarefa.',
   'Quase nunca serve a você, e entra em toda conversa com uma descrição longa. Vale desligar no '
   'claude.ai.',
   s='desligar')
_c('management-consultant', 'documentos',
   'Consultor de gestão (estilo grandes consultorias): estruturar problema, dimensionar mercado, '
   'montar entrega executiva.',
   'Problema de negócio ou de estratégia: fora do seu uso atual. Vale desligar no claude.ai.',
   s='desligar')
_c('google-workspace', 'documentos',
   'Cria e edita arquivos do Google (Docs, Planilhas, Apresentações) no seu Drive.',
   'Só se o trabalho for em arquivos do Google. Fora do seu uso atual: vale desligar no '
   'claude.ai.',
   s='desligar')
_c('import-memory', 'documentos',
   'Importa para a memória do Claude o que você exportou de outro assistente de IA.',
   'Uma vez só, se for migrar de outro assistente. Depois, vale desligar no claude.ai.',
   s='desligar')
_c('web-artifacts-builder', 'documentos',
   'Monta páginas web mais elaboradas (React, Tailwind) como artefato do claude.ai.',
   'Quase nunca aqui: o app é feito no Claude Code. Vale desligar no claude.ai.',
   s='desligar')
_c('llm-council', 'documentos',
   'Passa uma decisão por um “conselho” de 5 assessores de IA, que analisam, avaliam uns aos '
   'outros e dão um veredito.',
   'Decisão com risco e várias opções, quando você quer vê-la testada por ângulos diferentes. '
   'Decisão sua: ligar ou desligar.',
   '“Rode o conselho: devo fazer A ou B?”',
   s='decidir')
_c('prompt-master', 'documentos',
   'Escreve prompts melhores para outras ferramentas de IA (imagem, vídeo, agentes de código).',
   'Só quando você pede um prompt para outra ferramenta. Decisão sua: ligar ou desligar.',
   '“Escreva um prompt para gerar a imagem de um troféu.”',
   s='decidir')
_c('deep-research', 'documentos',
   'Pesquisa em várias fontes, com subagentes, e entrega um relatório com as fontes citadas.',
   'Pergunta que pede comparar opções, revisar literatura ou entender um assunto a fundo. Antes, '
   'procure na biblioteca `claude-kit\\pesquisas\\INDICE.md`; depois, salve lá.',
   '“Pesquise o que há de evidência sobre alta precoce em pediatria e me dê um relatório.”',
   s='decidir')
_c('built-in-browser', 'documentos',
   'O navegador de dentro do app Claude Desktop: abre páginas, lê e clica, com os seus logins.',
   'Só no app Desktop, quando a tarefa precisa de um site. No VS Code e no terminal não existe.',
   '“Abra a página do sistema e confira se o relatório de hoje saiu.”',
   s='decidir')
_c('chrome-browser', 'documentos',
   'O Claude no seu Chrome, pela extensão: usa a aba e os logins de verdade.',
   'Quando a tarefa precisa do seu Chrome e você quer acompanhar. Nunca em sistema com dado de '
   'paciente (ADR-0001).',
   '“No Chrome, baixe o PDF da norma que está aberta na aba.”',
   s='decidir')
_c('computer-use', 'documentos',
   'O Claude usa programas do computador: tira foto da tela, clica e digita, pelo app Desktop.',
   'Tarefa num programa sem outro jeito de automatizar. Nunca com dado de paciente na tela '
   '(ADR-0001).',
   '“Abra o Bloco de Notas e cole esta lista.”',
   s='decidir')
_c('find-skills', 'documentos',
   'Ajuda a achar e instalar skills novas quando você pergunta “tem skill para…?”.',
   'Quando você quer ampliar o que o Claude faz. Decisão sua: ligar ou desligar.',
   '“Existe alguma skill para criar flashcards do Anki?”',
   s='decidir')

# --- Do projeto -------------------------------------------------------
_c('desosp-app:revisar-tela', 'projeto',
   'Lista de conferência de uma tela antes do aceite: hierarquia, estados vazio e de erro, '
   'teclado, contraste, impressão e vocabulário, com a prova de cada ponto.',
   'Ao fechar toda janela que mudou tela, antes das fotos e do aceite. O `CLAUDE.md` do app '
   'manda usá-la.',
   '“Revise a tela de pendências antes do aceite.”')
_c('desosp-censo:rodada', 'projeto',
   'Roteiro da rodada do censo, na ordem certa: conferir a pasta `entrada`, extração, geração do '
   'censo e relatório.',
   'Antes de qualquer rodada, extração, reprocesso ou adendo. O `CLAUDE.md` do kernel manda '
   'carregá-la antes.',
   '“Processe a rodada da manhã.”')
# Os 12 comandos da planilha HC: o "para quê" é a primeira linha do próprio arquivo (@disco). A
# ordem e o "quando" seguem a seção "Fluxo de uma rodada" do CLAUDE.md do projeto (conferido em
# 06/10/2026).
_c('desosp-hc:nova-rodada', 'projeto', '@disco',
   'Passo 1 da rodada: abrir a rodada, com os arquivos das equipes e a master já baixados em '
   '`entrada`.')
_c('desosp-hc:rodada', 'projeto', '@disco',
   'Passo 2: corrigir e validar as unidades. Se a rodada não foi aberta, ele pede o '
   '[[desosp-hc:nova-rodada]] antes.')
_c('desosp-hc:consolidar', 'projeto', '@disco',
   'Passo 3: juntar tudo na master baixada no dia. Sem a master em `entrada`, ele para e pede.')
_c('desosp-hc:boletins', 'projeto', '@disco',
   'Passo 4: boletins e relatórios por hospital. Você aprova o texto de cada hospital antes.')
_c('desosp-hc:placar', 'projeto', '@disco',
   'Passo 4, só na rodada do mês: depois dos boletins e antes das mensagens.')
_c('desosp-hc:informativo', 'projeto', '@disco',
   'Passo 5, só quando muda a versão da planilha.')
_c('desosp-hc:mensagem', 'projeto', '@disco',
   'Passo 6: as mensagens de WhatsApp, uma para o grupo das equipes e outra para o gestor.')
_c('desosp-hc:ler-pdfs', 'projeto', '@disco',
   'Para conferir os PDFs da rodada (boletins, relatórios, informativo) lendo-os de volta.')
_c('desosp-hc:fechar-rodada', 'projeto', '@disco',
   'Passo 7, só depois que você aprovou a entrega. Mostra o plano e só executa com o seu “sim”.')
_c('desosp-hc:nova-versao', 'projeto', '@disco',
   'Quando a planilha ganha versão nova: discute as mudanças com você antes de construir.')
_c('desosp-hc:novo-hospital', 'projeto', '@disco',
   'Quando um hospital entra na rede.')
_c('desosp-hc:verificar', 'projeto', '@disco',
   'Antes de uma rodada, de uma versão nova ou de um registro importante.')


# ----------------------------------------------------------------------
# ler o que está instalado
# ----------------------------------------------------------------------
@dataclass
class Item:
    chave: str             # identidade única: 'grill-me', 'desosp-hc:rodada'
    nome: str              # o que aparece na coluna Skill (e o que se digita depois da barra)
    origem: str            # kit | plugin | sincronizada | projeto | claude-code
    rotulo: str            # a linha de origem, sob o nome
    tipo: str              # skill | comando
    como: str              # so_barra | auto | so_claude | comando
    descricao: str         # lida do disco ('' nas embutidas)
    projeto: str = ''
    grupo: str = ''        # nas do kit: o plugin (nucleo, planejamento, engenharia, kit)


def _conferir_caminho(caminho: Path) -> None:
    """Defesa extra: nenhum caminho que passe por pasta de dados chega a ser aberto."""
    ruins = {parte.lower() for parte in caminho.parts} & {n.lower() for n in PROIBIDAS}
    if ruins:
        raise PermissionError(f'recusei abrir {caminho}: passa por pasta de dados ({sorted(ruins)})')


def _desaspar(texto: str) -> str:
    texto = texto.strip()
    if len(texto) >= 2 and texto[0] == texto[-1] and texto[0] in '"\'':
        interior = texto[1:-1]
        if texto[0] == '"':
            return interior.replace('\\"', '"').replace('\\\\', '\\')
        return interior.replace("''", "'")
    return texto


def _ler_texto(caminho: Path) -> list[str]:
    _conferir_caminho(caminho)
    return caminho.read_text(encoding='utf-8-sig').splitlines()


def ler_frontmatter(linhas: list[str]) -> tuple[dict, list[str]]:
    """O bloco --- do começo do arquivo (só o primeiro nível) e o resto das linhas."""
    if not linhas or linhas[0].strip() != '---':
        return {}, linhas
    fim = None
    for i in range(1, len(linhas)):
        if linhas[i].strip() == '---':
            fim = i
            break
    if fim is None:
        return {}, linhas
    bloco, resto = linhas[1:fim], linhas[fim + 1:]
    dados: dict[str, str] = {}
    i = 0
    while i < len(bloco):
        m = re.match(r'^([A-Za-z0-9_-]+):\s*(.*)$', bloco[i])
        i += 1
        if not m:
            continue
        chave, valor = m.group(1), m.group(2).strip()
        continuacao = []
        while i < len(bloco) and (bloco[i].startswith((' ', '\t')) or not bloco[i].strip()):
            continuacao.append(bloco[i].strip())
            i += 1
        continuacao = [c for c in continuacao if c]
        if valor in ('|', '>', '|-', '>-', '|+', '>+'):
            texto = ' '.join(continuacao)
        elif not valor:
            texto = ''                                    # mapa aninhado (metadata:): ignora
        elif valor[0] in '"\'':
            texto = _desaspar(' '.join([valor] + continuacao))
        else:
            texto = ' '.join([valor] + continuacao)
        dados[chave] = texto
    return dados, resto


def _ler_md(f: Path):
    """Frontmatter e corpo de um .md. Se o caminho passa por pasta de dados: avisa e devolve None."""
    try:
        return ler_frontmatter(_ler_texto(f))
    except PermissionError as erro:
        AVISOS.append(str(erro))
        return None


def _verdade(valor) -> bool:
    return str(valor or '').strip().lower() in ('true', 'yes', '1')


def _como(fm: dict, comando: bool) -> str:
    if _verdade(fm.get('disable-model-invocation')):
        return 'so_barra'
    if str(fm.get('user-invocable', '')).strip().lower() == 'false':
        return 'so_claude'
    return 'comando' if comando else 'auto'


def _primeira_linha(linhas: list[str]) -> str:
    for l in linhas:
        l = l.strip().lstrip('#').strip()
        if l:
            return l
    return ''


def _json(caminho: Path, avisar: bool = True) -> dict:
    try:
        return json.loads(caminho.read_text(encoding='utf-8-sig'))
    except (OSError, ValueError) as erro:
        if avisar:
            AVISOS.append(f'não consegui ler {caminho.name} ({erro.__class__.__name__}): o guia '
                          'pode estar incompleto')
        return {}


def grupos_do_kit() -> list[tuple[str, Path]]:
    """(grupo, pasta do plugin), na ordem do .claude-plugin\\marketplace.json."""
    mkt = _json(KIT / '.claude-plugin' / 'marketplace.json')
    return [(p['name'], (KIT / p['source']).resolve()) for p in mkt.get('plugins', [])]


def ler_kit() -> list[Item]:
    itens = []
    for grupo, pasta in grupos_do_kit():
        base = pasta / 'skills'
        if not base.is_dir():
            continue
        for d in sorted(base.iterdir(), key=lambda x: x.name.lower()):
            if not d.is_dir() or d.name.startswith(('.', '_')):
                continue
            f = d / 'SKILL.md'
            if not f.is_file():
                AVISOS.append(f'pasta do kit sem SKILL.md: plugins\\{grupo}\\skills\\{d.name}')
                continue
            lido = _ler_md(f)
            if lido is None:
                continue
            fm, _ = lido
            nome = fm.get('name') or d.name
            itens.append(Item(nome, nome, 'kit', f'Kit · {grupo}', 'skill', _como(fm, False),
                              fm.get('description', ''), grupo=grupo))
    return itens


def conferir_ligacoes() -> tuple[int, int]:
    """O que o instalar_kit.py --verificar acha (junção sobrando, nome repetido, grupo no escopo errado...)."""
    import importlib.util
    spec = importlib.util.spec_from_file_location('instalar_kit', KIT / 'scripts' / 'instalar_kit.py')
    ik = importlib.util.module_from_spec(spec)
    sys.modules['instalar_kit'] = ik
    spec.loader.exec_module(ik)
    achados = ik.verificar(ik.Ambiente())
    for a in achados:
        AVISOS.append(f'instalar_kit.py --verificar: {a}')
    total = sum(len(s) for s in ik.plugins_do_kit(ik.Ambiente()).values())
    return total, len(achados)


def ler_ligacoes() -> tuple[list[str], list[list[str]]]:
    """A tabela projeto × grupo: o que vale em cada pasta (usuário < projeto < local, nos dois sentidos)."""
    usuario = _json(CLAUDE_HOME / 'settings.json', avisar=False).get('enabledPlugins') or {}
    grupos = [g for g, _ in grupos_do_kit()]

    def efetivo(pasta: Path) -> dict[str, bool]:
        ef = {k: bool(v) for k, v in usuario.items()}
        for nome in ('settings.json', 'settings.local.json'):
            f = pasta / '.claude' / nome
            _conferir_caminho(f)
            ef.update({k: bool(v) for k, v in (_json(f, avisar=False).get('enabledPlugins') or {}).items()})
        return ef

    pastas = [('raiz (C:\\CLAUDE-PROJETOS)', RAIZ_PROJETOS)]
    for d in sorted(RAIZ_PROJETOS.iterdir(), key=lambda x: x.name.lower()):
        n = d.name.lower()
        if (not d.is_dir() or n.startswith(('.', '_')) or n in {p.lower() for p in PROIBIDAS}
                or n.endswith(('-dados', '-backups'))):
            continue
        pastas.append((d.name, d))
    linhas = []
    for nome, pasta in pastas:
        ef = efetivo(pasta)
        linha = [nome] + ['sim' if ef.get(f'{g}@claude-kit') else '—' for g in grupos]
        outros = [k.split('@')[0] for k in sorted(ef) if ef[k] and not k.endswith('@claude-kit')]
        linhas.append(linha + [', '.join(outros) or '—'])
    return ['Pasta'] + grupos + ['Plugins de fora'], linhas


def _pasta_plugin(nome: str, mkt: str) -> Path | None:
    inst = _json(CLAUDE_HOME / 'plugins' / 'installed_plugins.json', avisar=False)
    for reg in (inst.get('plugins') or {}).get(f'{nome}@{mkt}', []):
        p = Path(reg.get('installPath', ''))
        if p.is_dir():
            return p
    base = CLAUDE_HOME / 'plugins' / 'cache' / mkt / nome
    if base.is_dir():
        versoes = sorted((d for d in base.iterdir() if d.is_dir()),
                         key=lambda d: d.stat().st_mtime, reverse=True)
        if versoes:
            return versoes[0]
    return None


def ler_plugins() -> tuple[list[Item], list[dict]]:
    cfg = _json(CLAUDE_HOME / 'settings.json')
    ligados = [k for k, v in (cfg.get('enabledPlugins') or {}).items() if v]
    itens, resumo = [], []
    for chave in ligados:
        nome, _, mkt = chave.partition('@')
        if mkt == 'claude-kit':          # os grupos do kit vêm do ler_kit, direto da pasta
            continue
        pasta = _pasta_plugin(nome, mkt)
        reg = dict(plugin=nome, mkt=mkt, versao='?', skills=0, comandos=0, achado=pasta is not None)
        if pasta is None:
            AVISOS.append(f'plugin ligado e não achado no cache: {chave}')
            resumo.append(reg)
            continue
        reg['versao'] = pasta.name
        rotulo = f'Plugin {nome}'
        for f in sorted((pasta / 'skills').glob('*/SKILL.md')):
            lido = _ler_md(f)
            if lido is None:
                continue
            fm, _ = lido
            n = fm.get('name') or f.parent.name
            itens.append(Item(n, n, 'plugin', rotulo, 'skill', _como(fm, False),
                              fm.get('description', ''), projeto=nome))
            reg['skills'] += 1
        for f in sorted((pasta / 'commands').glob('*.md')):
            lido = _ler_md(f)
            if lido is None:
                continue
            fm, resto = lido
            n = f.stem
            itens.append(Item(n, n, 'plugin', rotulo, 'comando', _como(fm, True),
                              fm.get('description') or _primeira_linha(resto), projeto=nome))
            reg['comandos'] += 1
        resumo.append(reg)
    return itens, resumo


def ler_sincronizadas() -> tuple[list[Item], str]:
    itens, data_manifesto = [], ''
    base = CLAUDE_HOME / 'skills' / 'synced'
    manifestos = sorted(base.glob('*/manifest.json')) if base.is_dir() else []
    vistos = set()
    for man in manifestos:
        dados = _json(man)
        data_manifesto = dt.datetime.fromtimestamp(man.stat().st_mtime).strftime('%d/%m/%Y')
        for s in dados.get('skills', []):
            nome = s.get('name', '')
            if not nome or nome in vistos:
                continue
            vistos.add(nome)
            fm = {}
            f = man.parent / nome / 'SKILL.md'
            if f.is_file():
                lido = _ler_md(f)
                fm = lido[0] if lido else {}
            else:
                AVISOS.append(f'skill sincronizada sem pasta no disco: {nome}')
            itens.append(Item(nome, nome, 'sincronizada', 'claude.ai (sincronizada)', 'skill',
                              _como(fm, False), s.get('description') or fm.get('description', '')))
    if not manifestos:
        AVISOS.append('nenhum manifest.json de skills sincronizadas foi achado')
    return itens, data_manifesto


def ler_projetos() -> tuple[list[Item], list[str]]:
    itens, lidos = [], []
    pastas = dict(PROJETOS)
    # projetos novos: qualquer pasta irmã com .claude\skills ou .claude\commands (só olha isso)
    for d in sorted(RAIZ_PROJETOS.iterdir(), key=lambda x: x.name.lower()):
        n = d.name.lower()
        if (not d.is_dir() or d.name in pastas or d.resolve() == KIT or n.startswith(('.', '_'))
                or n in {p.lower() for p in PROIBIDAS} or n.endswith(('-dados', '-backups'))):
            continue
        if (d / '.claude' / 'skills').is_dir() or (d / '.claude' / 'commands').is_dir():
            pastas[d.name] = d.name
            AVISOS.append(f'projeto com skills ou comandos que o gerador não conhecia: {d.name}')
    for pasta, apelido in pastas.items():
        raiz = RAIZ_PROJETOS / pasta / '.claude'
        rotulo = f'Projeto: {apelido}'
        achou = False
        if (raiz / 'skills').is_dir():
            for f in sorted((raiz / 'skills').glob('*/SKILL.md')):
                lido = _ler_md(f)
                if lido is None:
                    continue
                fm, _ = lido
                n = fm.get('name') or f.parent.name
                itens.append(Item(f'{pasta}:{n}', n, 'projeto', rotulo, 'skill', _como(fm, False),
                                  fm.get('description', ''), projeto=pasta))
                achou = True
        if (raiz / 'commands').is_dir():
            for f in sorted((raiz / 'commands').glob('*.md')):
                lido = _ler_md(f)
                if lido is None:
                    continue
                fm, resto = lido
                n = f.stem
                itens.append(Item(f'{pasta}:{n}', n, 'projeto', rotulo, 'comando', _como(fm, True),
                                  fm.get('description') or _primeira_linha(resto), projeto=pasta))
                achou = True
        if achou:
            lidos.append(pasta)
    return itens, lidos


def ler_nativas() -> list[Item]:
    return [Item(n, n, 'claude-code', 'Claude Code (embutida)', 'skill', 'auto', '')
            for n in NATIVAS]


@dataclass
class Inventario:
    itens: list[Item]
    kit_ligadas: int
    kit_soltas: int
    plugins: list[dict]
    data_manifesto: str
    projetos_lidos: list[str]

    def por_chave(self) -> dict[str, Item]:
        return {i.chave: i for i in self.itens}

    def de(self, origem: str) -> list[Item]:
        return [i for i in self.itens if i.origem == origem]


def inventariar() -> Inventario:
    kit = ler_kit()
    ligadas, soltas = conferir_ligacoes()
    plug, resumo = ler_plugins()
    sinc, data_man = ler_sincronizadas()
    proj, lidos = ler_projetos()
    itens = kit + plug + sinc + proj + ler_nativas()
    pessoais = CLAUDE_HOME / 'commands'
    if pessoais.is_dir() and any(pessoais.glob('*.md')):
        AVISOS.append(r'há comandos pessoais em ~\.claude\commands que este guia ainda não lê')
    repetidos = {i.chave for i in itens if sum(1 for j in itens if j.chave == i.chave) > 1}
    for r in sorted(repetidos):
        AVISOS.append(f'nome repetido entre origens diferentes: {r} (a tabela mostra os dois)')
    return Inventario(itens, ligadas, soltas, resumo, data_man, lidos)


def plural(n: int, um: str, varios: str) -> str:
    return f'{n} {um if n == 1 else varios}'


def milhar(n: int) -> str:
    return f'{n:,}'.replace(',', '.')


def tokens_estimados(i: Item) -> int:
    """Estimativa grosseira do que a descrição custa em toda conversa: caracteres ÷ 4."""
    return (len(i.nome) + min(len(i.descricao), 1536)) // 4


# ----------------------------------------------------------------------
# cores, estilos e utilitários de Word (o jeito do gerar_guia_rodada.py)
# ----------------------------------------------------------------------
AZUL = RGBColor(0x1F, 0x38, 0x64)
CINZA = RGBColor(0x59, 0x59, 0x59)
VERMELHO = RGBColor(0xB0, 0x2A, 0x1F)
AMBAR = RGBColor(0x9C, 0x5B, 0x00)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
F_AZUL, F_AZUL_CLARO, F_CINZA = '1F3864', 'DEEAF6', 'F2F2F2'
F_ALERTA, F_OK, F_AMBAR = 'FDE9E7', 'E8F3E8', 'FFF4D6'
LARGURA = 17.0            # cm úteis (A4 com margem de 2 cm)
doc: Document = None      # type: ignore[assignment]
CONTAGEM = {'catalogo': 0, 'sem_curadoria': 0}
_marcadores = [0]

TITULOS = {
    1: 'Como ler este guia',
    2: 'Os quatro jeitos de uma skill entrar no trabalho',
    3: 'Catálogo por grupo',
    4: 'O fluxo do grill em 6 passos',
    5: 'Skills por projeto — proposta (o plano de 06/10 confirma)',
    6: 'O que ficou de fora e por quê',
    7: 'Como manter',
}


def _estilos(data: str) -> None:
    st = doc.styles
    st['Normal'].font.name = 'Calibri'
    st['Normal'].font.size = Pt(11)
    st['Normal'].paragraph_format.space_after = Pt(4)
    for nome, tam, antes in (('Title', 26, 0), ('Heading 1', 16, 18), ('Heading 2', 13, 12),
                             ('Heading 3', 11, 8)):
        s = st[nome]
        s.font.name = 'Calibri'
        s.font.size = Pt(tam)
        s.font.bold = True
        s.font.color.rgb = AZUL
        s.paragraph_format.space_before = Pt(antes)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    # fontes: tira as marcas de "fonte do tema", que mandam mais que o nome da fonte
    rfonts = [s.element.get_or_add_rPr().find(qn('w:rFonts')) for s in
              (st['Normal'], st['Title'], st['Heading 1'], st['Heading 2'], st['Heading 3'])]
    padrao = st.element.find(qn('w:docDefaults'))
    if padrao is not None:
        rpr = padrao.find(qn('w:rPrDefault'))
        if rpr is not None and rpr.find(qn('w:rPr')) is not None:
            rf = rpr.find(qn('w:rPr')).find(qn('w:rFonts'))
            if rf is not None:
                rfonts.append(rf)
            lang = rpr.find(qn('w:rPr')).find(qn('w:lang'))
            if lang is not None:
                lang.set(qn('w:val'), 'pt-BR')
                lang.set(qn('w:eastAsia'), 'pt-BR')
    for rf in rfonts:
        if rf is None:
            continue
        for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rf.get(qn(k)) is not None:
                del rf.attrib[qn(k)]
        for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(k), 'Calibri')
    # bordas das tabelas em cinza claro (a "Table Grid" vem em preto)
    grade = st['Table Grid'].element.find(qn('w:tblPr'))
    bordas = grade.find(qn('w:tblBorders')) if grade is not None else None
    if bordas is not None:
        for b in bordas:
            b.set(qn('w:color'), 'BFBFBF')
            b.set(qn('w:sz'), '4')
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    rod = sec.footer.paragraphs[0]
    rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(rod, f'Guia das skills do Claude Code · gerado em {data} · página ', cor=CINZA, tam=8)
    _campo(rod, 'PAGE')
    _run(rod, ' de ', cor=CINZA, tam=8)
    _campo(rod, 'NUMPAGES')
    cfg = doc.settings.element
    for cs in cfg.iter(qn('w:compatSetting')):
        if cs.get(qn('w:name')) == 'compatibilityMode':
            cs.set(qn('w:val'), '15')               # Word 2013 ou mais novo: sem o aviso de compatibilidade
    tema = cfg.find(qn('w:themeFontLang'))
    if tema is not None:
        tema.set(qn('w:val'), 'pt-BR')
    # o modelo do python-docx traz uma miniatura em branco; sem ela, o Explorador mostra o ícone do Word
    pacote = doc.part.package
    for rid, rel in list(pacote.rels.items()):
        if rel.reltype.endswith('/thumbnail'):
            pacote.rels.pop(rid)
    cp = doc.core_properties
    cp.title = 'Guia das skills do Claude Code'
    cp.subject = 'O que cada skill faz, quando usar e como chamar'
    cp.author = 'claude-kit (gerado por script)'
    cp.last_modified_by = 'claude-kit'
    cp.comments = f'Gerado por {AQUI.name}. Para mudar, edite o script e rode de novo.'
    cp.language = 'pt-BR'
    agora = dt.datetime.now()
    cp.created = cp.modified = agora


def _run(par, texto, negrito=False, cor=None, tam=None, mono=False, italico=False):
    r = par.add_run(texto)
    r.bold, r.italic = negrito, italico
    if cor is not None:
        r.font.color.rgb = cor
    if tam:
        r.font.size = Pt(tam)
    if mono:
        r.font.name = 'Consolas'
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Consolas')
    return r


def _campo(par, codigo: str) -> None:
    def fld(tipo):
        r = par.add_run()
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), tipo)
        r._r.append(fc)
        return r
    r1 = fld('begin')
    r2 = par.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = f' {codigo} '
    r2._r.append(it)
    r3 = fld('separate')
    r4 = par.add_run('1')
    r5 = fld('end')
    for r in (r1, r2, r3, r4, r5):
        r.font.size = Pt(8)
        r.font.color.rgb = CINZA


def _rico(par, texto, **kw):
    """**negrito** e `código` (monoespaçado, azul) dentro do texto."""
    base_negrito = kw.pop('negrito', False)
    for pedaco in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', texto):
        if not pedaco:
            continue
        if pedaco.startswith('**'):
            _run(par, pedaco[2:-2], negrito=True, **kw)
        elif pedaco.startswith('`'):
            _run(par, pedaco[1:-1], mono=True, cor=AZUL, negrito=base_negrito,
                 **{k: v for k, v in kw.items() if k == 'tam'})
        else:
            _run(par, pedaco, negrito=base_negrito, **kw)
    return par


def p(texto, **kw):
    manter = kw.pop('manter', False)
    par = _rico(doc.add_paragraph(), texto, **kw)
    if manter:
        par.paragraph_format.keep_with_next = True
    return par


def h1(n: int, nova_pagina=False):
    par = doc.add_heading(f'{n}. {TITULOS[n]}', level=1)
    par.paragraph_format.page_break_before = nova_pagina
    _marcar(par, f'sec{n}')
    return par


def h2(texto):
    return doc.add_heading(texto, level=2)


def h3(texto):
    return doc.add_heading(texto, level=3)


def itens_lista(lista):
    for t in lista:
        _rico(doc.add_paragraph(style='List Bullet'), t)


def _marcar(par, nome: str) -> None:
    _marcadores[0] += 1
    ini = OxmlElement('w:bookmarkStart')
    ini.set(qn('w:id'), str(_marcadores[0]))
    ini.set(qn('w:name'), nome)
    fim = OxmlElement('w:bookmarkEnd')
    fim.set(qn('w:id'), str(_marcadores[0]))
    ppr = par._p.find(qn('w:pPr'))
    if ppr is not None:
        ppr.addnext(ini)
    else:
        par._p.insert(0, ini)
    par._p.append(fim)


def _link(par, texto: str, ancora: str) -> None:
    h = OxmlElement('w:hyperlink')
    h.set(qn('w:anchor'), ancora)
    h.set(qn('w:history'), '1')
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    cor = OxmlElement('w:color')
    cor.set(qn('w:val'), '1F3864')
    rpr.append(cor)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = texto
    r.append(t)
    h.append(r)
    par._p.append(h)


_APOS_SHD = ('w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark')
_APOS_MAR = ('w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark')


def _inserir_tcpr(tc_pr, novo, sucessores):
    """Põe `novo` antes do primeiro irmão que deve vir depois dele (ordem do esquema do Word)."""
    for nome in sucessores:
        achado = tc_pr.find(qn(nome))
        if achado is not None:
            achado.addprevious(novo)
            return
    tc_pr.append(novo)


def _sombra(cel, cor):
    tc_pr = cel._element.get_or_add_tcPr()
    for velho in tc_pr.findall(qn('w:shd')):
        tc_pr.remove(velho)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), cor)
    _inserir_tcpr(tc_pr, shd, _APOS_SHD)


def _margens(cel, cm=0.18, vertical=None):
    """Margem interna da célula; `vertical` (topo e base) pode ser menor que a dos lados."""
    tc_pr = cel._element.get_or_add_tcPr()
    for velho in tc_pr.findall(qn('w:tcMar')):
        tc_pr.remove(velho)
    mar = OxmlElement('w:tcMar')
    for lado in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{lado}')
        valor = vertical if (vertical is not None and lado in ('top', 'bottom')) else cm
        e.set(qn('w:w'), str(int(valor * 567)))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    _inserir_tcpr(tc_pr, mar, _APOS_MAR)


def _larguras(tab, cms, cabecalho=False):
    tab.autofit = False
    for gc, w in zip(tab._tbl.tblGrid.findall(qn('w:gridCol')), cms):
        gc.set(qn('w:w'), str(int(w * 567)))
    for i, linha in enumerate(tab.rows):
        tr_pr = linha._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement('w:cantSplit'))
        if cabecalho and i == 0:
            tr_pr.append(OxmlElement('w:tblHeader'))
            for cel in linha.cells:
                for par in cel.paragraphs:
                    par.paragraph_format.keep_with_next = True
        for cel, w in zip(linha.cells, cms):
            cel.width = Cm(w)


def _texto_cel(cel, texto, negrito=False, cor=None, tam=10, centro=False, italico=False):
    par = cel.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    if centro:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, linha in enumerate(texto.split('\n')):
        if i:
            par = cel.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            if centro:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _rico(par, linha, negrito=negrito, cor=cor, tam=tam, italico=italico)


def _paragrafo_cel(cel, antes=0):
    par = cel.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.space_before = Pt(antes)
    return par


def espaco(pt=6):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    par.paragraph_format.line_spacing = Pt(pt)


def _juntar(tab):
    """Prende as linhas umas às outras: o Word mantém a tabela inteira na mesma página, se couber."""
    for linha in tab.rows[:-1]:
        for cel in linha.cells:
            for par in cel.paragraphs:
                par.paragraph_format.keep_with_next = True


def tabela(cabecalho, linhas, cms, tam=9.5, destaque=True, fundos=None, juntar=False):
    """Tabela comum. `linhas`: listas de texto (com **negrito** e `código`). `fundos`: {linha: cor}.
    `juntar`: mantém a tabela inteira na mesma página (só para tabelas curtas)."""
    tab = doc.add_table(rows=1 + len(linhas), cols=len(cabecalho))
    tab.style = 'Table Grid'
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, t in enumerate(cabecalho):
        c = tab.rows[0].cells[j]
        _sombra(c, F_AZUL)
        _margens(c)
        _texto_cel(c, t, negrito=True, cor=BRANCO, tam=tam)
    for i, linha in enumerate(linhas, start=1):
        for j, t in enumerate(linha):
            c = tab.rows[i].cells[j]
            if fundos and (i - 1) in fundos:
                _sombra(c, fundos[i - 1])
            elif i % 2 == 0:
                _sombra(c, F_CINZA)
            _margens(c)
            _texto_cel(c, t, tam=tam, negrito=(destaque and j == 0))
    _larguras(tab, cms, cabecalho=True)
    if juntar:
        _juntar(tab)
    espaco()
    return tab


def quadro(titulo, linhas, cor=F_AZUL_CLARO):
    """Caixa de destaque: uma célula sombreada, título em negrito e linhas de texto."""
    tab = doc.add_table(rows=1, cols=1)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tab.rows[0].cells[0]
    _sombra(c, cor)
    _margens(c, 0.3)
    _texto_cel(c, titulo, negrito=True, cor=AZUL, tam=11)
    for l in linhas:
        par = c.add_paragraph()
        par.paragraph_format.space_after = Pt(2)
        _rico(par, l, tam=10)
    _larguras(tab, [LARGURA])
    espaco()


def codigo(linhas):
    tab = doc.add_table(rows=1, cols=1)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tab.rows[0].cells[0]
    _sombra(c, F_CINZA)
    _margens(c, 0.3)
    par = c.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    for i, l in enumerate(linhas):
        if i:
            par = c.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
        _run(par, l, mono=True, tam=9.5)
    _larguras(tab, [LARGURA])
    espaco()


def caixas(blocos, cor=F_AZUL_CLARO):
    """Esquema em linha SEM setas: caixas lado a lado. bloco = (título, corpo, rodapé)."""
    n = len(blocos)
    vao = 0.3
    w = (LARGURA - vao * (n - 1)) / n
    tab = doc.add_table(rows=1, cols=2 * n - 1)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    cms = []
    for i in range(2 * n - 1):
        c = tab.rows[0].cells[i]
        if i % 2 == 0:
            titulo, corpo, rodape = blocos[i // 2]
            _sombra(c, cor)
            _margens(c, 0.2)
            _texto_cel(c, titulo, negrito=True, cor=AZUL, tam=10.5)
            par = _paragrafo_cel(c, 3)
            _rico(par, corpo, tam=9.5)
            if rodape:
                par = _paragrafo_cel(c, 4)
                _rico(par, rodape, tam=8.5, cor=CINZA, italico=True)
            cms.append(w)
        else:
            cms.append(vao)
    _larguras(tab, cms)
    espaco()


def faixa(etapas, cor=F_AZUL_CLARO):
    """Uma linha de caixas pequenas ligadas por setas: etapa → etapa → etapa."""
    n = len(etapas)
    seta = 0.7
    w = (LARGURA - seta * (n - 1)) / n
    tab = doc.add_table(rows=1, cols=2 * n - 1)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    cms = []
    for i in range(2 * n - 1):
        c = tab.rows[0].cells[i]
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if i % 2 == 0:
            numero, titulo = etapas[i // 2]
            _sombra(c, cor)
            _margens(c, 0.12)
            _texto_cel(c, numero, negrito=True, cor=AZUL, tam=13, centro=True)
            par = _paragrafo_cel(c)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(par, titulo, negrito=True, cor=AZUL, tam=8.5)
            cms.append(w)
        else:
            _margens(c, 0.0)
            _texto_cel(c, '→', negrito=True, cor=AZUL, tam=14, centro=True)
            cms.append(seta)
    _larguras(tab, cms)
    espaco()


def cartao_numerado(numero, titulo, linhas):
    """Uma linha: o número grande num bloco azul e, ao lado, o título e o texto do passo."""
    tab = doc.add_table(rows=1, cols=2)
    tab.style = 'Table Grid'
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = tab.rows[0].cells
    a.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _sombra(a, F_AZUL)
    _margens(a, 0.1)
    _texto_cel(a, str(numero), negrito=True, cor=BRANCO, tam=20, centro=True)
    _margens(b, 0.25)
    _texto_cel(b, titulo, negrito=True, cor=AZUL, tam=11.5)
    for l in linhas:
        par = _paragrafo_cel(b, 2)
        _rico(par, l, tam=10)
    _larguras(tab, [1.5, LARGURA - 1.5])
    espaco(4)


def _inicio_bloco() -> int:
    """Posição em que o próximo elemento entra no corpo (o sectPr fica sempre por último)."""
    return len(doc.element.body) - 1


def _prender_bloco(inicio: int) -> None:
    """Prende uns aos outros todos os elementos escritos desde `inicio`, menos o último: o Word
    mantém o bloco inteiro na mesma página, se couber (senão, passa tudo para a página seguinte)."""
    corpo = list(doc.element.body)
    fim = len(corpo) - 1
    for el in corpo[inicio:fim - 1]:
        pars = [el] if el.tag == qn('w:p') else list(el.iter(qn('w:p')))
        for par in pars:
            Paragraph(par, doc._body).paragraph_format.keep_with_next = True


# ----------------------------------------------------------------------
# texto: referências [[skill]] e rótulos do "Como chamar"
# ----------------------------------------------------------------------
INV: dict[str, Item] = {}
CITADAS: set[str] = set()
FALTANDO: set[str] = set()

ROTULOS = {      # como -> (rótulo, fundo)
    'so_barra': ('Só por /', F_AZUL_CLARO),
    'auto': ('Sozinho ou /', F_OK),
    'comando': ('Por /', F_AZUL_CLARO),
    'so_claude': ('Só sozinho', F_CINZA),
}


def _ref(m) -> str:
    chave = m.group(1)
    CITADAS.add(chave)
    item = INV.get(chave)
    nome = chave.split(':')[-1]
    if item is None:
        FALTANDO.add(chave)
        return f'`{nome}` (NÃO INSTALADA)'
    if item.como in ('so_barra', 'comando'):
        return f'`/{nome}`'
    return f'`{nome}`'


def expandir(texto: str) -> str:
    return re.sub(r'\[\[([^\]]+)\]\]', _ref, texto)


def _pode_quebrar(texto: str) -> str:
    """Põe um espaço invisível depois de / e _ para o Word poder quebrar caminhos longos."""
    return texto.replace('/', '/\u200b').replace('_', '_\u200b')


def _cortar(texto: str, n: int = 190) -> str:
    """A descrição lida do disco, limpa: a primeira frase (se couber) e no máximo n letras."""
    texto = re.sub(r'\s+', ' ', texto.replace('$ARGUMENTS', '')).strip().rstrip(':').strip()
    primeira = re.split(r'(?<=[a-zà-ú\)])\.\s+(?=[A-ZÀ-Ú])', texto, maxsplit=1)[0]
    if 30 <= len(primeira) < len(texto):
        texto = primeira if primeira.endswith('.') else primeira + '.'
    if len(texto) <= n:
        return texto
    return texto[:n].rsplit(' ', 1)[0].rstrip(',;:.') + '…'


def _digita(i: Item) -> str:
    return '(só o Claude)' if i.como == 'so_claude' else f'/{i.nome}'


# ----------------------------------------------------------------------
# o catálogo: uma tabela por grupo
# ----------------------------------------------------------------------
def _cor_marca(t: str):
    return VERMELHO if t.startswith('vale desligar') else AMBAR


def linha_catalogo(tab_linha, item: Item, cur: dict | None, tam=9):
    c0, c1, c2, c3 = tab_linha.cells
    for c in (c0, c1, c2, c3):
        _margens(c, 0.14, vertical=0.07)
    # 1) nome, origem e marcas
    par = c0.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    _run(par, item.nome, negrito=True, cor=AZUL, tam=tam, mono=True)
    par = _paragrafo_cel(c0)
    _run(par, item.rotulo, cor=CINZA, tam=7.5)
    marcas = []
    if cur and cur.get('s') == 'desligar':
        marcas.append('vale desligar no claude.ai')
        _sombra(c0, F_ALERTA)
    elif cur and cur.get('s') == 'decidir':
        marcas.append('decisão sua: ligar ou desligar')
        _sombra(c0, F_AMBAR)
    if cur and cur.get('t'):
        marcas.append(cur['t'])
    for m in marcas:
        par = _paragrafo_cel(c0, 1)
        _run(par, m, negrito=True, cor=_cor_marca(m), tam=7.5)
    # 2) como chamar (vem do arquivo da skill, não deste texto)
    rotulo, fundo = ROTULOS[item.como]
    _sombra(c1, fundo)
    par = c1.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    _run(par, rotulo, negrito=True, cor=AZUL, tam=9)
    par = _paragrafo_cel(c1, 1)
    _run(par, _digita(item), mono=True, cor=AZUL, tam=7.5)
    # 3) para quê e 4) quando usar
    if cur is None:
        return
    if cur['p'] == '@disco':
        texto_p = _pode_quebrar(_cortar(item.descricao))
    else:
        texto_p = expandir(cur['p'])
    _texto_cel(c2, texto_p, tam=tam)
    _texto_cel(c3, expandir(cur['q']), tam=tam)
    if cur.get('e'):
        par = _paragrafo_cel(c3, 3)
        ex = cur['e']
        _run(par, 'Ex.: ', negrito=True, cor=CINZA, tam=tam - 1)
        if ex.startswith('/'):
            _run(par, ex, mono=True, cor=AZUL, tam=tam - 1.5)
        else:
            _run(par, ex, italico=True, cor=CINZA, tam=tam - 0.5)


def tabela_catalogo(pares):
    cab = ['Skill', 'Como chamar', 'Para quê', 'Quando usar']
    cms = [3.5, 3.1, 4.9, 5.5]
    tab = doc.add_table(rows=1 + len(pares), cols=4)
    tab.style = 'Table Grid'
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, t in enumerate(cab):
        c = tab.rows[0].cells[j]
        _sombra(c, F_AZUL)
        _margens(c, 0.14)
        _texto_cel(c, t, negrito=True, cor=BRANCO, tam=9)
    for i, (item, cur) in enumerate(pares, start=1):
        linha_catalogo(tab.rows[i], item, cur)
        CONTAGEM['catalogo'] += 1
    _larguras(tab, cms, cabecalho=True)
    espaco()


def tabela_sem_curadoria(itens):
    cab = ['Skill', 'Origem', 'Como chamar', 'O que diz o arquivo (lido do disco)']
    cms = [3.5, 3.4, 2.9, 7.2]
    tab = doc.add_table(rows=1 + len(itens), cols=4)
    tab.style = 'Table Grid'
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, t in enumerate(cab):
        c = tab.rows[0].cells[j]
        _sombra(c, F_AZUL)
        _margens(c, 0.14)
        _texto_cel(c, t, negrito=True, cor=BRANCO, tam=9)
    for i, item in enumerate(itens, start=1):
        c0, c1, c2, c3 = tab.rows[i].cells
        for c in (c0, c1, c2, c3):
            _margens(c, 0.14, vertical=0.07)
        _run(c0.paragraphs[0], item.nome, negrito=True, cor=AZUL, tam=9, mono=True)
        c0.paragraphs[0].paragraph_format.space_after = Pt(0)
        _texto_cel(c1, item.rotulo, tam=8.5, cor=CINZA)
        rotulo, fundo = ROTULOS[item.como]
        _sombra(c2, fundo)
        _texto_cel(c2, rotulo + '\n`' + _digita(item) + '`', tam=8.5)
        _texto_cel(c3, _pode_quebrar(_cortar(item.descricao, 230)) or '(sem descrição no arquivo)',
                   tam=8.5)
        CONTAGEM['sem_curadoria'] += 1
    _larguras(tab, cms, cabecalho=True)
    espaco()


def grupo_de(item: Item) -> str | None:
    """Skill do kit: o grupo do plugin dela. As outras: o grupo do CURADO."""
    cur = CURADO.get(item.chave)
    if cur is None:
        return None
    return item.grupo if item.origem == 'kit' and item.grupo else cur['g']


def agrupar(inv: Inventario):
    ordem = {c: n for n, c in enumerate(CURADO)}
    por_grupo: dict[str, list[tuple[Item, dict]]] = {g[0]: [] for g in GRUPOS}
    sem: list[Item] = []
    for item in inv.itens:
        g = grupo_de(item)
        if g is None or g not in por_grupo:
            sem.append(item)
        else:
            por_grupo[g].append((item, CURADO[item.chave]))
    for g in por_grupo.values():
        g.sort(key=lambda par: ordem[par[0].chave])
    return por_grupo, sem


# ----------------------------------------------------------------------
# as seções
# ----------------------------------------------------------------------
def capa(inv: Inventario, agora: dt.datetime) -> None:
    doc.add_paragraph('Guia das skills do Claude Code', style='Title')
    p('O que cada skill faz, quando usar e como chamar.', cor=CINZA)
    p(f'Gerado em {agora:%d/%m/%Y} às {agora:%H:%M}, a partir do que está instalado neste PC.',
      cor=CINZA, tam=9)
    quadro('Este guia se regenera', [
        'Mudou uma skill, um plugin ou um projeto? Rode o comando abaixo e o Word é refeito:'],
        cor=F_AZUL_CLARO)
    codigo([COMANDO])
    p('O texto vive no script; o Word é gerado. Para mudar uma frase, edite o script (o dicionário '
      '`CURADO`) e rode de novo. O gerador lê o que está instalado de verdade: uma skill nova, '
      'sem texto próprio, aparece em “Sem descrição curada”, no fim da seção 3.')

    h2('Neste momento, instalado neste PC')
    kit, plug = inv.de('kit'), inv.de('plugin')
    sinc, proj, nat = inv.de('sincronizada'), inv.de('projeto'), inv.de('claude-code')
    n_sb = sum(1 for i in kit if i.como == 'so_barra')
    nomes_plug = ', '.join(f"{r['plugin']}" for r in inv.plugins) or '—'
    tabela(['Origem', 'Itens', 'Detalhe', 'Onde está'], [
        ['Kit (grupos)', str(len(kit)),
         f'{n_sb} só `/` · {len(kit) - n_sb} que o Claude chama sozinho · '
         + ', '.join(f'{g} ({sum(1 for i in kit if i.grupo == g)})' for g, _ in grupos_do_kit()),
         f'`{KIT / "plugins"}\\<grupo>\\skills`'],
        ['Plugins ligados', str(len(plug)),
         f'{plural(len(inv.plugins), "plugin", "plugins")} ({nomes_plug}): '
         f'{plural(sum(1 for i in plug if i.tipo == "skill"), "skill", "skills")} e '
         f'{plural(sum(1 for i in plug if i.tipo == "comando"), "comando", "comandos")}',
         f'`{CLAUDE_HOME / "plugins" / "cache"}`'],
        ['claude.ai (sincronizadas)', str(len(sinc)),
         'valem em todos os projetos' + (f'; lista de {inv.data_manifesto}' if inv.data_manifesto
                                         else ''),
         f'`{CLAUDE_HOME / "skills" / "synced"}`'],
        ['Projetos', str(len(proj)),
         f'{plural(sum(1 for i in proj if i.tipo == "skill"), "skill", "skills")} e '
         f'{plural(sum(1 for i in proj if i.tipo == "comando"), "comando", "comandos")}, em '
         + ', '.join(f'`{x}`' for x in inv.projetos_lidos),
         r'`<projeto>\.claude\skills` e `\.claude\commands`'],
        ['Claude Code (embutidas)', str(len(nat)), 'lista fixa dentro do gerador (não dá para ler '
         'do disco)', 'dentro do programa'],
        ['Total', str(len(inv.itens)), '', ''],
    ], [3.4, 1.2, 5.0, 7.4], tam=9, juntar=True)

    h2('Neste guia')
    par = doc.add_paragraph(style='List Bullet')
    _run(par, 'Antes de tudo: ', negrito=True, cor=AZUL)
    _link(par, TITULO_OPERACAO, 'operacao')
    for n in range(1, 8):
        par = doc.add_paragraph(style='List Bullet')
        _run(par, f'{n}. ', negrito=True, cor=AZUL)
        _link(par, TITULOS[n], f'sec{n}')


OPERACAO = KIT / 'COMO_OPERAR.md'
TITULO_OPERACAO = 'Como você opera (até o F3)'


def parte_operacao() -> None:
    """A 1ª parte do guia: o COMO_OPERAR.md do kit, em Word (títulos ##, listas, numeradas, tabelas)."""
    par = doc.add_heading(TITULO_OPERACAO, level=1)
    par.paragraph_format.page_break_before = True
    _marcar(par, 'operacao')
    if not OPERACAO.is_file():
        AVISOS.append(f'{OPERACAO.name} não existe: a parte da operação ficou vazia')
        return
    linhas = OPERACAO.read_text(encoding='utf-8').splitlines()
    texto: list[str] = []
    marcas: list[str] = []

    def soltar():
        if texto:
            p(' '.join(t.strip() for t in texto))
            texto.clear()
        if marcas:
            itens_lista(marcas[:])
            marcas.clear()

    i = 0
    while i < len(linhas):
        l = linhas[i]
        if l.startswith('# '):
            i += 1
            continue
        if l.startswith('## '):
            soltar()
            h2(l[3:].strip())
        elif l.startswith('|'):
            soltar()
            bloco = []
            while i < len(linhas) and linhas[i].startswith('|'):
                bloco.append([c.strip() for c in linhas[i].strip().strip('|').split('|')])
                i += 1
            cab, corpo = bloco[0], [r for r in bloco[1:] if not set(''.join(r)) <= set('-: ')]
            w = LARGURA / len(cab)
            tabela(cab, corpo, [w * 0.8] + [(LARGURA - w * 0.8) / (len(cab) - 1)] * (len(cab) - 1)
                   if len(cab) > 1 else [LARGURA], tam=9.5, destaque=False)
            continue
        elif re.match(r'^\d+\. ', l):
            soltar()
            n, resto = l.split('. ', 1)
            p(f'**{n}.** {resto.strip()}')
        elif re.match(r'^\s*- ', l):
            if texto:
                p(' '.join(t.strip() for t in texto))
                texto.clear()
            marcas.append(l.split('- ', 1)[1].strip())
        elif not l.strip():
            soltar()
        elif marcas and l.startswith('  ') and not l.startswith('   ' + ' '):
            marcas[-1] += ' ' + l.strip()
        else:
            if marcas:
                itens_lista(marcas[:])
                marcas.clear()
            texto.append(l)
        i += 1
    soltar()


def sec1(inv: Inventario) -> None:
    h1(1, nova_pagina=True)
    p('Este guia responde, para cada skill, a duas perguntas: **o que ela faz** e **como ela entra '
      'no trabalho**: você chama com `/`, o Claude chama sozinho, o projeto manda usar, ou um '
      'documento espera o assunto aparecer (seção 2).')
    quadro('Em 30 segundos', [
        'Quer saber se precisa digitar `/`? Veja a coluna “Como chamar”, na seção 3.',
        'Quer saber em que momento usar cada skill? Seções 4 e 5.',
        'Instalou ou removeu uma skill? Seção 7 (e regenere o guia).'])
    h2('A coluna “Como chamar” (seção 3)')
    p('Ela é lida do arquivo de cada skill (o campo `disable-model-invocation`), não deste texto. '
      'Por isso nunca desatualiza.', manter=True)
    tab_marcas = tabela(['Marca', 'O que quer dizer'], [
        ['Só por /', 'Você digita `/nome`. O Claude não a chama sozinho e a descrição dela nem '
                     'entra na conversa: **não gasta contexto até você chamar**.'],
        ['Sozinho ou /', 'O Claude a carrega quando o seu pedido combina com a descrição; você '
                         'também pode digitar `/nome`. A **descrição fica em toda conversa**.'],
        ['Por /', 'Comando (um arquivo na pasta `commands` de um projeto ou de um plugin): você '
                  'digita `/nome`.'],
        ['Só sozinho', 'Aparece só para o Claude: não está no menu `/`. (Nenhuma hoje.)'],
    ], [3.4, 13.6], fundos=None, tam=9.5, juntar=True)
    # pinta a primeira coluna com a cor de cada marca (a mesma das tabelas do catálogo)
    for i, chave in enumerate(('so_barra', 'auto', 'comando', 'so_claude'), start=1):
        _sombra(tab_marcas.rows[i].cells[0], ROTULOS[chave][1])
    h2('A coluna “Skill”, três linhas')
    itens_lista([
        '**O nome**, como se digita depois da barra. Skill de plugin e do claude.ai também aceita '
        'o nome completo (`plugin:nome`, por exemplo `anthropic-skills:docx`); o curto basta '
        'enquanto nenhum outro comando usar o mesmo nome.',
        '**A origem**, em cinza: Kit (pessoal), Plugin, claude.ai, Projeto ou Claude Code. Ela diz '
        'onde a skill vale (seção 2).',
        '**As marcas**, em cor: “experimental”, “em reavaliação”, “vale desligar no claude.ai”, '
        '“decisão sua”.'])
    h2('O que vem do disco e o que foi escrito à mão')
    n_hc = sum(1 for i in inv.itens if i.projeto == 'desosp-hc')
    n_nat = len(inv.de('claude-code'))
    itens_lista([
        '**Do disco, sempre atual:** o nome, a descrição, a marca “só /”, a origem e a lista do '
        'que existe. Skill instalada aparece; skill removida some.',
        '**À mão, no dicionário do gerador:** o grupo, o “para quê”, o “quando usar” e o exemplo. '
        'Skill instalada sem texto à mão cai em “Sem descrição curada”, no fim da seção 3: é o '
        'aviso de que algo mudou.',
        f'**Nos {n_hc} comandos da planilha HC**, o “para quê” é a primeira linha do próprio '
        'arquivo do comando, lida do disco.',
        f'**As {n_nat} embutidas do Claude Code** não existem no disco: o nome e a marca vêm de '
        'uma lista fixa dentro do gerador, conferida em 06/10/2026.'])
    h2('Palavras usadas aqui')
    tabela(['Palavra', 'O que quer dizer'], [
        ['Skill', 'um roteiro escrito que o Claude carrega para fazer bem um tipo de tarefa.'],
        ['Contexto', 'tudo o que o Claude tem diante de si numa conversa. Quanto mais cheio, mais '
                     'caro e mais lento. É o custo que mais pesa.'],
        ['Token', 'o pedaço de texto em que o contexto e o uso são contados.'],
        ['Modelo e esforço', 'o modelo é qual Claude responde (Opus, Sonnet); o esforço (`/effort`) é o '
                             'quanto ele pensa antes de responder. Mais esforço custa mais.'],
        ['`CLAUDE.md`', 'o arquivo de regras que o Claude lê ao abrir a janela. O pessoal vale em '
                        'todos os projetos; o do projeto, só nele.'],
        ['Gancho (hook)', 'programa que roda sozinho num momento fixo, como ao abrir a janela ou '
                          'antes de cada comando. Não é skill.'],
        ['Plugin', 'pacote que traz skills, ganchos e comandos prontos. Liga e desliga por '
                   'projeto.'],
        ['MCP', 'conector que liga o Claude a outro programa ou serviço.'],
        ['Subagente', 'ajudante que trabalha numa conversa à parte e devolve só o resumo, sem '
                      'encher o seu contexto.'],
        ['Handoff', 'passar o trabalho de uma janela para a próxima, com o que ela precisa saber.'],
        ['Especificação e tarefas', 'o que construir e por quê; e os pedaços pequenos em que isso '
                                    'se divide (tickets).'],
        ['Git e commit', 'o git guarda o histórico das versões do código; commit é salvar uma versão, '
                         'com uma legenda.'],
        ['ADR', 'uma decisão registrada num arquivo curto, com o motivo (fica em `docs/adr/`).'],
    ], [3.8, 13.2], tam=9.5, juntar=True)


def sec2(inv: Inventario) -> None:
    h1(2)
    p('Uma skill pode entrar no seu trabalho de quatro jeitos. Saber qual é o jeito de cada uma '
      'responde à pergunta “eu preciso lembrar de chamar?”.', manter=True)
    tabela(['Jeito', 'Como acontece', 'O que custa', 'Exemplo daqui'], [
        ['**a) Você chama** com `/nome`',
         'Você digita `/` e o nome da skill na caixa de texto. As marcadas **só /** rodam apenas '
         'assim.',
         'Nada até você chamar: a descrição das “só /” nem entra na conversa.',
         '`/grill-me`, `/to-spec`, `/retro`'],
        ['**b) O Claude chama sozinho**',
         'Quando o seu pedido combina com a descrição da skill, o Claude a carrega. Você não '
         'digita nada.',
         'A descrição fica em toda conversa: 70 a 150 tokens por skill.',
         'Você diz “tem um bug difícil no merge” e entra a `diagnosing-bugs`.'],
        ['**c) Programada no projeto**',
         'O `CLAUDE.md` do projeto manda usar a skill num momento certo (“ao fechar a janela, use '
         'a de fechamento”). Ninguém precisa lembrar.',
         'Só uma linha no `CLAUDE.md`.',
         expandir('Decidido em 06/10; ainda a implantar, de modo geral, em cada projeto. O começo '
                  'já existe: o `CLAUDE.md` pessoal manda usar a [[uso-do-claude]] e a '
                  '[[recursos-do-projeto]]; o do kernel manda carregar a [[desosp-censo:rodada]] '
                  'antes de toda rodada; o do app manda usar a [[desosp-app:revisar-tela]] ao '
                  'fechar janela de tela.')],
        ['**d) Documento sob demanda**',
         r'Regras em `.claude\rules` com `paths:`, e documentos que só carregam quando o assunto '
         'aparece. Não são skills.',
         'Só quando o Claude abre um arquivo daquele caminho (ou o documento).',
         'No kernel, a regra `home-care.md` só carrega quando o Claude mexe nos arquivos de home '
         'care.'],
    ], [3.0, 5.0, 3.5, 5.5], tam=9, juntar=True)

    caixas([
        ('a · Você chama', '`/nome` na caixa de texto.', 'Quem aciona: você'),
        ('b · O Claude chama', 'O pedido combina com a descrição da skill.',
         'Quem aciona: o Claude'),
        ('c · O projeto manda', 'O `CLAUDE.md` diz em que momento usar.', 'Quem aciona: o projeto'),
        ('d · O documento espera', 'Carrega quando o assunto (o arquivo) aparece.',
         'Quem aciona: o assunto'),
    ])

    auto = [i for i in inv.itens if i.como == 'auto' and i.origem in ('kit', 'plugin', 'sincronizada')]
    total = sum(tokens_estimados(i) for i in auto)
    so_barra = sum(1 for i in inv.de('kit') if i.como == 'so_barra')
    decididas = [i for i in inv.itens
                 if (CURADO.get(i.chave) or {}).get('s') in ('desligar', 'decidir')]
    tok_dec = sum(tokens_estimados(i) for i in decididas)
    quadro('A conta de hoje (estimativa: caracteres da descrição ÷ 4)', [
        f'O Claude pode chamar sozinho **{len(auto)} skills** (kit, plugins e claude.ai). As '
        f'descrições delas somam **cerca de {milhar(total)} tokens em toda conversa**, fora as '
        'embutidas do Claude Code.',
        f'As {so_barra} skills do kit marcadas “só /” não entram nessa conta.',
        f'As {len(decididas)} skills do claude.ai marcadas “vale desligar” ou “decisão sua” (grupo '
        f'Documentos, seção 3) respondem por **cerca de {milhar(tok_dec)} tokens** desse total.',
    ], cor=F_AMBAR)

    h2('Onde cada skill vale')
    p('Uma skill pessoal vale em **todas as janelas de todos os projetos** desta máquina. Uma skill '
      'de projeto vale **só ali**. Um plugin **liga e desliga por projeto**.', manter=True)
    caixas([
        ('Skill pessoal', 'Todas as janelas, de todos os projetos desta máquina. Não se desliga '
                          'por projeto.', 'Hoje: só as do claude.ai; as do kit vêm por plugin'),
        ('Skill de projeto', 'Só dentro daquele projeto: aparece quando a janela abre naquela '
                             'pasta.', 'Hoje: ' + (', '.join(PROJETOS.get(x, x) for x in inv.projetos_lidos)
                                                   or 'nenhum')),
        ('Plugin', 'Traz skills, ganchos e comandos juntos. Liga e desliga por projeto '
                   '(`enabledPlugins`).', 'Hoje: ' + (', '.join(r['plugin'] for r in inv.plugins)
                                                       or 'nenhum')),
    ])
    tabela(['Tipo', 'Onde mora', 'Como se liga ou desliga'], [
        ['Skill do kit (plugin)', f'`{KIT / "plugins"}\\<grupo>\\skills`, carregada da própria pasta '
                                  '(o kit é um marketplace local)',
         'Por grupo, no `enabledPlugins` do usuário (nucleo, planejamento) ou do projeto (engenharia, '
         'sdd, kit). Depois de `git pull`: `instalar_kit.py --verificar`; de terceiros: '
         '`atualizar_terceiros.py`.'],
        ['Skill de projeto', r'`<projeto>\.claude\skills` e `<projeto>\.claude\commands`',
         'Só existe naquele projeto.'],
        ['Plugin de fora', f'`{CLAUDE_HOME / "plugins" / "cache"}`',
         '`enabledPlugins` no `settings.json` do usuário ou do projeto.'],
        ['claude.ai (sincronizada)', f'`{CLAUDE_HOME / "skills" / "synced"}` (não mexer)',
         'No próprio claude.ai: o Claude Code recebe a lista pronta.'],
        ['Claude Code (embutida)', 'dentro do programa', 'Já vem; não se instala nem se desliga.'],
    ], [3.8, 7.2, 6.0], tam=9, juntar=True)

    h2('O que NÃO é skill')
    tabela(['O que é', 'Para quê', 'Exemplo daqui'], [
        ['**Gancho (hook)**',
         'Programa que roda sozinho num momento fixo (abrir a janela, antes ou depois de cada '
         'comando). Não gasta contexto e você não chama.',
         'O `cc-safety-net` (antes de cada comando), o de abertura do app (ao abrir a janela) e o '
         'das cores de pasta (depois de cada comando).'],
        ['**Plugin**',
         'Pacote que traz skills, ganchos e comandos juntos. A skill do plugin é só uma parte.',
         'O `cc-safety-net` traz um gancho e uma skill.'],
        ['**MCP**',
         'Conector que liga o Claude a outro programa ou serviço (um site, um banco, o Notion).',
         'O conector do Claude Docs, que cria e edita documentos no claude.ai.'],
    ], [3.3, 7.4, 6.3], tam=9, juntar=True)


def sec3(inv: Inventario) -> None:
    h1(3)
    p('Uma tabela por grupo. Em cada linha: o nome da skill, **como chamar** (lido do arquivo dela), '
      'para que serve e quando usar, com um exemplo do que pedir. A ordem dentro do grupo é a ordem '
      'em que costumam aparecer no trabalho.')
    p('Cor da coluna “Como chamar”: **azul**, você digita `/nome`; **verde**, o Claude chama '
      'sozinho (e você também pode digitar `/nome`).', cor=CINZA, tam=9.5)
    por_grupo, sem = agrupar(inv)
    notas = {
        'interface': ('A `revisar-tela`, que confere uma tela do app antes do aceite, só existe no '
                      'app e está no grupo “Do projeto”.'
                      if 'desosp-app:revisar-tela' in INV else None),
        'projeto': ('Em cada linha, a origem diz a qual projeto ela pertence: app é `desosp-app`, '
                    'kernel é `desosp-censo` e planilha HC é `desosp-hc`. A palavra `rodada` existe '
                    'em dois: no kernel é uma skill; na planilha HC, um comando. Cada janela só '
                    'enxerga a do seu projeto.'),
        'documentos': (r'Não desligue apagando arquivos em `skills\synced`: o claude.ai sincroniza '
                       'essa pasta sozinho.'),
    }
    for gid, titulo, intro in GRUPOS:
        pares = por_grupo[gid]
        if not pares:
            continue
        h2(titulo)
        p(expandir(intro), manter=True)
        if notas.get(gid):
            p(expandir(notas[gid]), manter=True)
        tabela_catalogo(pares)
        if gid == 'controles' and any(i.chave == 'cc-safety-net' for i, _ in pares):
            quadro('O que o gancho do cc-safety-net barra', [
                '`git reset --hard`, `git push --force`, `git clean -f`, `git checkout --`, apagar '
                'arquivos fora da pasta de trabalho e ler `.env` e arquivos de credenciais, '
                'no Bash e no PowerShell.',
                '**Lacuna conhecida:** apagar uma pasta DENTRO da pasta do projeto não é barrado.',
                r'Há uma política proposta em `scripts\safety-net-politica.json`, que protege também '
                'no terminal as pastas de dados do app. Quem a aplica é você, não o Claude: o plugin '
                'não deixa o Claude mudar a própria política.'],
                cor=F_AMBAR)
        if gid == 'projeto':
            p('Nos comandos da planilha HC, a ordem e o “quando usar” seguem a seção “Fluxo de uma '
              'rodada” do `CLAUDE.md` do projeto; o “para quê” é a primeira linha do arquivo do '
              'comando.', cor=CINZA, tam=9)
    h2('Sem descrição curada — atualize o gerador')
    if sem:
        quadro('Atenção: há skills instaladas sem texto no dicionário', [
            f'Foram achadas {len(sem)} skills ou comandos que o dicionário `CURADO` do gerador '
            'ainda não conhece. O guia só mostra o que o disco diz sobre elas. Acrescente-as ao '
            f'dicionário (grupo, para quê, quando usar, exemplo) em `{AQUI}` e rode de novo.'],
            cor=F_ALERTA)
        tabela_sem_curadoria(sem)
    else:
        quadro('Nenhuma: tudo o que está instalado tem descrição curada', [
            'Conferido na geração deste guia. Se uma skill nova for instalada sem texto no '
            'dicionário, ela aparece aqui.'], cor=F_OK)


def sec4(inv: Inventario) -> None:
    inicio = _inicio_bloco()
    h1(4)
    p('Para uma ideia nova ou uma mudança grande. É a ordem em que as skills da seção 3 se '
      'encadeiam.', manter=True)
    faixa([('1', 'Entrevista'), ('2', 'Ver antes'), ('3', 'Especificar'), ('4', 'Tarefas'),
           ('5', 'Implementar'), ('6', 'Fechar')])
    passos = [
        ('Entrevista',
         ['[[grill-me]] (ou [[grill-with-docs]], se o assunto tem vocabulário próprio ou a decisão '
          'precisa ficar escrita).',
          '**Modelo:** Opus 5.5 com `/effort high`, com o modo plano desligado.']),
        ('Ver antes de decidir',
         ['Dúvida que só se responde vendo: [[prototype]].',
          'Resposta que depende de outra pessoa: [[to-questionnaire]].']),
        ('Especificação',
         ['[[to-spec]], na mesma janela da entrevista.',
          'Projeto novo: antes dele, a [[recursos-do-projeto]].']),
        ('Tarefas',
         ['[[to-tickets]] divide a especificação em tarefas, cada uma dizendo do que depende.',
          'Grande demais para uma janela: [[wayfinder]].']),
        ('Implementar',
         ['[[implement]], uma tarefa por vez, com o teste primeiro.',
          '[[implement-spec]] só quando as tarefas são independentes entre si (é caro).']),
        ('Fechar',
         ['Passar o trabalho (handoff): a [[uso-do-claude]] orienta o prompt da próxima janela; a '
          '[[claude-handoff]] passa o trabalho a uma sessão nova, ainda experimental.',
          'E a retrospectiva: [[retro]].']),
    ]
    for n, (titulo, linhas) in enumerate(passos, start=1):
        cartao_numerado(n, titulo, [expandir(l) for l in linhas])
    quadro('Mudança pequena? Pule para o passo 5, ou nem use skill.', [
        'Corrigir uma palavra, ajustar uma cor, mudar um texto: peça direto ao Claude. O fluxo '
        'inteiro é para o que tem decisão no meio.'])
    _prender_bloco(inicio)


PROPOSTA = [
    ('App (desosp-app)', [
        ('Pensar uma funcionalidade nova', ['grill-with-docs'], ' + '),
        ('Virar trabalho', ['to-spec', 'to-tickets', 'implement'], ' → '),
        ('Código novo ou correção, com teste', ['tdd'], ' + '),
        ('Bug difícil', ['diagnosing-bugs'], ' + '),
        ('Dúvida que só se responde vendo', ['prototype'], ' + '),
        ('Tela nova', ['frontend-design', 'web-design-guidelines'], ' + '),
        ('Antes do aceite de uma tela', ['desosp-app:revisar-tela'], ' + '),
    ]),
    ('Kernel (desosp-censo)', [
        ('Pensar uma regra ou funcionalidade', ['grill-with-docs', 'domain-modeling'], ' + '),
        ('Testar regras sobre datas e âncoras', ['property-based-testing'], ' + '),
        ('Conferir o código contra o documento que o especifica', ['spec-to-code-compliance'],
         ' + '),
        ('Quando o código ficar difícil de mexer', ['improve-codebase-architecture'], ' + '),
        ('Toda rodada do censo', ['desosp-censo:rodada'], ' + '),
    ]),
    ('Planilha HC (desosp-hc)', [
        ('Mexer em planilhas, PDFs e documentos', ['xlsx', 'pdf', 'docx'], ' + '),
        ('Decisão que depende de outra pessoa', ['to-questionnaire'], ' + '),
        ('Conferir o código contra o documento', ['spec-to-code-compliance'], ' + '),
        ('A rodada de atualização, na ordem do `CLAUDE.md` do projeto', ['@hc_rodada'], ' → '),
        ('Quando for preciso', ['@hc_extra'], ' · '),
    ]),
]
# os comandos da planilha HC, na ordem da rodada (os que não estiverem aqui caem em "quando for
# preciso", para um comando novo nunca ficar de fora)
HC_RODADA = ['nova-rodada', 'rodada', 'consolidar', 'boletins', 'placar', 'informativo',
             'mensagem', 'fechar-rodada']


def _como_entra(itens: list[Item]) -> str:
    barra = [i for i in itens if i.como in ('so_barra', 'comando')]
    auto = [i for i in itens if i.como not in ('so_barra', 'comando')]
    fmt = lambda xs: ', '.join(f'`{"/" if x.como in ("so_barra", "comando") else ""}{x.nome}`'
                               for x in xs)
    if barra and not auto:
        t = 'você chama com `/`'
    elif auto and not barra:
        t = 'o Claude chama sozinho'
    else:
        t = f'você chama {fmt(barra)}; o Claude chama {fmt(auto)}'
    if any(i.origem == 'projeto' for i in itens):
        t += '; já existe no projeto'
    return t


def sec5(inv: Inventario) -> None:
    h1(5, nova_pagina=True)
    h2('Que grupo cada pasta liga hoje (lido dos settings, não deste texto)')
    p('“sim” = o grupo aparece numa janela aberta naquela pasta. Vale o escopo de usuário '
      '(`~\\.claude\\settings.json`), e o `.claude\\settings.json` (e o `.local`) da pasta vence. '
      'O grupo engenharia só liga num projeto de código, na fase dele; até lá, o projeto fica sem ele.',
      cor=CINZA, tam=9)
    cab, linhas = ler_ligacoes()
    largura_g = 2.0
    tabela(cab, linhas, [4.4] + [largura_g] * (len(cab) - 2) + [LARGURA - 4.4 - largura_g * (len(cab) - 2)],
           tam=9, destaque=False)
    h2('Em que momento chamar cada skill (proposta)')
    citadas = {c for _, linhas in PROPOSTA for _, chaves, _ in linhas for c in chaves
               if not c.startswith('@')}
    faltam = sorted(c.split(':')[-1] for c in citadas if c not in INV)
    instaladas = ('As skills abaixo já estão instaladas (seção 3).' if not faltam else
                  'ATENÇÃO: nem todas as skills abaixo estão instaladas; faltam: '
                  + ', '.join(f'`{n}`' for n in faltam) + '.')
    quadro('PROPOSTA: o plano de 06/10 confirma ou muda', [
        instaladas + ' O que ainda falta é escrever, no `CLAUDE.md` de cada projeto, em que '
        'momento chamar cada uma: uma seção curta “Skills deste projeto” (até 10 linhas) e, no '
        '`settings.json` do projeto, ligar só o plugin que serve ali (`enabledPlugins`). Decidido '
        'em 06/10; ainda a implantar.'], cor=F_ALERTA if faltam else F_AMBAR)
    p('Leitura: nomes com barra você digita; nomes sem barra o Claude chama sozinho (e você também '
      'pode digitar).', cor=CINZA, tam=9)
    cmds_hc = {i.nome: i for i in inv.itens if i.projeto == 'desosp-hc'}
    for projeto, linhas in PROPOSTA:
        h2(projeto)
        corpo = []
        for momento, chaves, sep in linhas:
            if chaves[0].startswith('@hc'):
                if chaves == ['@hc_rodada']:
                    itens = [cmds_hc[n] for n in HC_RODADA if n in cmds_hc]
                else:
                    itens = [i for n, i in cmds_hc.items() if n not in HC_RODADA]
                for i in itens:
                    CITADAS.add(i.chave)
                if not itens:
                    FALTANDO.add('desosp-hc:(comandos)')
                    continue
            else:
                itens = []
                for ch in chaves:
                    CITADAS.add(ch)
                    if ch in INV:
                        itens.append(INV[ch])
                    else:
                        FALTANDO.add(ch)
                if not itens:
                    continue
            nomes = sep.join(f'`{"/" if i.como in ("so_barra", "comando") else ""}{i.nome}`'
                             for i in itens)
            corpo.append([momento, nomes, _como_entra(itens)])
        tabela(['Quando', 'Skill', 'Como entra'], corpo, [5.2, 6.4, 5.4], tam=9, destaque=False,
               juntar=True)
    h2('Projetos novos, ainda sem skills')
    tabela(['Projeto', 'Situação'], [
        ['Pesquisa clínica', 'definidas no grill de cada um'],
        ['Comunicação', 'definidas no grill de cada um'],
        ['Estudo para a residência', 'definidas no grill de cada um'],
        ['Produtividade pessoal', 'definidas no grill de cada um'],
    ], [6.0, 11.0], tam=9.5, juntar=True)


def sec6() -> None:
    h1(6)
    p('Avaliados e deixados de fora, com o motivo. O que está aqui não deve ser instalado sem uma '
      'nova decisão.', manter=True)
    tabela(['Item', 'Por que ficou de fora'], [
        ['claude-mem e claude-remember',
         'Gravam a conversa (inclusive o que o Claude lê), podem pegar dado de paciente e gastam a '
         'cota da assinatura.'],
        ['planning-with-files e Beads',
         'Criam um segundo dono do estado do trabalho, ao lado do `ESTADO.md` e do git.'],
        ['hookify', 'Os ganchos dele chamam `python3`, que neste Windows é o atalho da Microsoft '
                    'Store e falha.'],
        ['security-guidance', 'Chama o modelo a cada fim de turno (custo). O `/security-review` do '
                              'Claude Code já cobre.'],
        ['modern-web-guidance', 'Obrigatória em toda tarefa web, e roda `npx` sem versão fixa a '
                                'cada uso.'],
        ['ui-ux-pro-max', 'Repete o que já existe: `frontend-design` e `revisar-tela`.'],
        ['pyright-lsp', 'O VS Code já entrega os diagnósticos de código ao Claude.'],
        ['Superpowers, ECC, gstack, GSD e ruflo, inteiros',
         'Decidem o processo de trabalho e brigam com o método da casa (a família do grill, o '
         '`ESTADO.md`).'],
        ['Cartographer, Claude Code Setup e Headroom',
         'Avaliados em 06/10: não instalar. O Cartographer lê o código inteiro (dado solto iria '
         'para o mapa) e edita o `CLAUDE.md`; o Claude Code Setup sugere conectores já recusados; '
         'o Headroom é um intermediário que vê tudo em texto puro e, na extensão do VS Code, '
         'quebra a tela.'],
        ['oh-my-claudecode',
         'Só numa pasta de teste, com dado de mentira. **NUNCA rodar** o `/omc-setup`: ele '
         'sobrescreve o `CLAUDE.md` pessoal.'],
    ], [5.0, 12.0], tam=9.5, fundos={9: F_ALERTA})


def sec7(inv: Inventario) -> None:
    h1(7)
    h2('As skills do kit: conferir, atualizar, acrescentar')
    codigo([rf'python {KIT / "scripts" / "instalar_kit.py"} --verificar',
            rf'python {KIT / "scripts" / "atualizar_terceiros.py"} --todas',
            rf'python {KIT / "scripts" / "atualizar_terceiros.py"} <skill> --aplicar'])
    p('As skills moram nos plugins do kit (`claude-kit\\plugins\\<grupo>\\skills`) e carregam de lá. '
      'O `instalar_kit.py --verificar` acusa junção antiga, nome repetido e grupo no escopo errado (sem '
      'o `--verificar`, conserta). As de terceiros se atualizam pelo `atualizar_terceiros.py`, que '
      'baixa numa pasta de preparo, compara e, com `--aplicar`, troca a pasta e o `origem.json`. '
      'Nunca `npx skills add` direto: ele grava em `~\\.claude\\skills` e repete o nome.')
    h2('Plugins')
    marketplaces = sorted({r['mkt'] for r in inv.plugins})
    cmds = [f'claude plugin marketplace update {m}' for m in marketplaces]
    cmds += [f"claude plugin update {r['plugin']}@{r['mkt']}" for r in inv.plugins]
    cmds += ['claude plugin details <nome>      (mostra o custo em contexto)']
    codigo(cmds)
    if inv.plugins:
        tabela(['Plugin', 'Marketplace', 'Versão', 'Traz'],
               [[f"`{r['plugin']}`", r['mkt'], r['versao'],
                 f"{plural(r['skills'], 'skill', 'skills')} · "
                 f"{plural(r['comandos'], 'comando', 'comandos')}" if r['achado']
                 else 'NÃO ACHADO no cache'] for r in inv.plugins],
               [4.0, 5.0, 3.6, 4.4], tam=9, juntar=True)
    h2('Skills do claude.ai')
    p('Sincronizam sozinhas. Para desligar uma, use o próprio claude.ai; não mexa em '
      r'`skills\synced`.')
    h2('Custo e uso')
    itens_lista([
        '`claude plugin details <nome>` mostra o custo de contexto de um plugin.',
        '`/skill-doctor`, dentro do Claude Code, mostra o custo e o uso de cada skill.',
        'Instalou ou atualizou com a janela aberta? Abra uma janela nova para ter certeza de que '
        'o Claude Code enxerga a mudança.'])
    h2('Regenerar este guia')
    codigo([COMANDO])
    p('Quando o guia mostrar skill em “Sem descrição curada”, edite o dicionário `CURADO` do '
      f'gerador (`{AQUI}`), acrescentando grupo, para quê, quando usar e exemplo, e rode de novo. '
      'Mudou o uso de uma skill num projeto? Atualize também a proposta da seção 5. O Claude Code '
      'ganhou uma skill embutida nova? Acrescente o nome em `NATIVAS` e o texto no `CURADO`.')
    h2('O que o gerador conferiu nesta geração')
    linhas, fundos = conferencias(inv)
    tabela(['Conferência', 'Resultado'], linhas, [8.0, 9.0], tam=9, fundos=fundos, juntar=True)


def conferencias(inv: Inventario):
    linhas, fundos = [], {}

    def add(titulo, ok, texto):
        fundos[len(linhas)] = F_OK if ok else F_ALERTA
        linhas.append([titulo, ('OK: ' if ok else 'ATENÇÃO: ') + texto])

    kit = inv.de('kit')
    add('O kit instalado neste PC (`instalar_kit.py --verificar`)', inv.kit_soltas == 0,
        f'{len(kit)} skills em {len(grupos_do_kit())} grupos; 0 achados.'
        if inv.kit_soltas == 0 else f'{inv.kit_soltas} achado(s): rode `instalar_kit.py` (sem o '
        '`--verificar`, ele conserta); a lista está em “Outros avisos”.')
    nao_achados = [r['plugin'] for r in inv.plugins if not r['achado']]
    add('Plugins ligados achados no cache', not nao_achados,
        f'{len(inv.plugins)} plugins ligados, todos achados.' if not nao_achados
        else 'não achados: ' + ', '.join(nao_achados))
    add('Skills do claude.ai', bool(inv.de('sincronizada')),
        f"{len(inv.de('sincronizada'))} no manifesto." if inv.de('sincronizada')
        else 'manifesto não achado ou vazio.')
    add('Projetos lidos', bool(inv.projetos_lidos),
        ', '.join(inv.projetos_lidos) + r'. Só se leu `.claude\skills` e `.claude\commands` de '
        'cada um; nenhuma pasta de dados foi aberta.')
    sem = [i.nome for i in inv.itens if i.chave not in CURADO]
    add('Skills instaladas sem descrição curada', not sem,
        'nenhuma.' if not sem else f'{len(sem)}: ' + ', '.join(sem))
    sobra = [c for c in CURADO if c not in INV]
    if sobra:
        add('Texto do dicionário sem skill instalada', True,
            f'{len(sobra)} (não aparecem no catálogo): ' + ', '.join(sobra))
    add('Skills citadas nas seções 4 e 5 e nos textos', not FALTANDO,
        f'{len(CITADAS)} citadas, todas instaladas.' if not FALTANDO
        else 'não instaladas: ' + ', '.join(sorted(FALTANDO)))
    esperado = len([i for i in inv.itens if grupo_de(i) in {g[0] for g in GRUPOS}])
    add('Linhas no catálogo × skills instaladas',
        CONTAGEM['catalogo'] == esperado and CONTAGEM['catalogo'] + CONTAGEM['sem_curadoria']
        == len(inv.itens),
        f"{CONTAGEM['catalogo']} no catálogo + {CONTAGEM['sem_curadoria']} sem descrição curada "
        f'= {len(inv.itens)} instaladas.')
    outros = [a for a in AVISOS if not a.startswith(('skill do kit NÃO', 'plugin ligado'))]
    if outros:
        add('Outros avisos', False, ' | '.join(outros))
    return linhas, fundos


# ----------------------------------------------------------------------
# montar, salvar e conferir
# ----------------------------------------------------------------------
def construir(inv: Inventario, agora: dt.datetime) -> None:
    global doc
    doc = Document()
    INV.clear()
    INV.update(inv.por_chave())
    CITADAS.clear()
    FALTANDO.clear()
    CONTAGEM.update(catalogo=0, sem_curadoria=0)
    _estilos(f'{agora:%d/%m/%Y}')
    capa(inv, agora)
    parte_operacao()
    sec1(inv)
    sec2(inv)
    sec3(inv)
    sec4(inv)
    sec5(inv)
    sec6()
    sec7(inv)


def salvar(destino: Path) -> Path:
    """Grava num arquivo temporário e troca. Se o Word estiver com o .docx aberto, não força."""
    tmp = destino.with_name(f'~gerando_{destino.name}')
    doc.save(tmp)
    try:
        os.replace(tmp, destino)
        return destino
    except PermissionError:
        alt = destino.with_name(f'{destino.stem}_{dt.datetime.now():%Y-%m-%d_%H%M}{destino.suffix}')
        os.replace(tmp, alt)
        print(f'ATENÇÃO: {destino.name} está aberto (no Word?) e não foi mexido. '
              f'Gravei com outro nome: {alt}')
        return alt


def conferir_docx(arquivo: Path, inv: Inventario) -> bool:
    """Abre o .docx gravado e confere: cada skill instalada aparece uma vez, os títulos na ordem."""
    d = Document(arquivo)
    titulos = [par.text for par in d.paragraphs if par.style.name == 'Heading 1']
    esperados = [TITULO_OPERACAO] + [f'{n}. {TITULOS[n]}' for n in range(1, 8)]
    linhas = []
    for t in d.tables:
        cab = [c.text.strip() for c in t.rows[0].cells]
        if cab[:2] in (['Skill', 'Como chamar'], ['Skill', 'Origem']):
            for r in t.rows[1:]:
                pars = r.cells[0].paragraphs
                origem = pars[1].text.strip() if cab[1] == 'Como chamar' else r.cells[1].text.strip()
                linhas.append((pars[0].text.strip(), origem))
    esperado = sorted((i.nome, i.rotulo) for i in inv.itens)
    ok_titulos = titulos == esperados
    ok_linhas = sorted(linhas) == esperado
    repetidas = len(linhas) - len(set(linhas))
    print(f'Conferência do arquivo: {len(d.tables)} tabelas; títulos na ordem: '
          f'{"sim" if ok_titulos else "NÃO"}; linhas de skill: {len(linhas)} para '
          f'{len(inv.itens)} instaladas; repetidas: {repetidas}; '
          f'{"batem" if ok_linhas else "NÃO BATEM"}.')
    return ok_titulos and ok_linhas and repetidas == 0


def main() -> int:
    ap = argparse.ArgumentParser(description='Gera o guia das skills em Word.')
    ap.add_argument('--saida', type=Path, default=DESTINO, help='onde gravar o .docx')
    args = ap.parse_args()
    agora = dt.datetime.now()
    inv = inventariar()
    construir(inv, agora)
    arquivo = salvar(args.saida)
    ok = conferir_docx(arquivo, inv)
    kit, plug = inv.de('kit'), inv.de('plugin')
    sem = [i.nome for i in inv.itens if i.chave not in CURADO]
    print(f'Kit: {len(kit)} ({sum(1 for i in kit if i.como == "so_barra")} só /) · '
          f'plugins: {len(inv.plugins)} ligados, {len(plug)} itens · '
          f'claude.ai: {len(inv.de("sincronizada"))} · projetos: {len(inv.de("projeto"))} · '
          f'Claude Code: {len(inv.de("claude-code"))} · total: {len(inv.itens)}')
    print('Sem descrição curada: ' + (', '.join(sem) if sem else 'nenhuma'))
    sobra = [c for c in CURADO if c not in INV]
    if sobra:
        print('Texto do dicionário sem skill instalada (não aparece): ' + ', '.join(sobra))
    if FALTANDO:
        print('Citadas e NÃO instaladas: ' + ', '.join(sorted(FALTANDO)))
    for a in AVISOS:
        print('AVISO:', a)
    print(arquivo)
    return 0 if ok else 1


if __name__ == '__main__':
    raise SystemExit(main())
