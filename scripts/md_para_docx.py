# -*- coding: utf-8 -*-
"""
md_para_docx.py: converte um documento Markdown do kit em Word (.docx), para o Ettore ler.

Cobre o que os documentos do kit usam: títulos (#, ##, ###), parágrafos com linhas quebradas, listas com marcador
(com níveis) e numeradas, tabelas, **negrito**, *itálico* e `código`. As fontes entre parênteses que citam arquivo
(".md" ou "§") saem em cinza e menores; nos títulos, saem do texto (ficam no .md).

Uso: python scripts/md_para_docx.py docs/ARQUIVO.md [docs/ARQUIVO.docx]
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

FONTE = re.compile(r'(\((?:[^()]*?(?:\.md|§)[^()]*?)\))')
INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)')
ITEM = re.compile(r'^(\s*)([-*]|\d+\.)\s+(.*)$')
CINZA = RGBColor(0x80, 0x80, 0x80)


def corre(par, texto, negrito=False, italico=False, fonte=False):
    for parte in INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith('**') and parte.endswith('**') and len(parte) > 4:
            corre(par, parte[2:-2], True, italico, fonte)
        elif parte.startswith('`') and parte.endswith('`') and len(parte) > 2:
            r = par.add_run(parte[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(8 if fonte else 9.5)
            r.bold, r.italic = negrito, italico
            if fonte:
                r.font.color.rgb = CINZA
        elif parte.startswith('*') and parte.endswith('*') and len(parte) > 2:
            corre(par, parte[1:-1], negrito, True, fonte)
        else:
            r = par.add_run(parte)
            r.bold, r.italic = negrito, italico
            if fonte:
                r.font.size = Pt(8)
                r.font.color.rgb = CINZA


def texto_rico(par, texto):
    for parte in FONTE.split(texto):
        if parte:
            corre(par, parte, fonte=bool(FONTE.fullmatch(parte)))


def tabela(doc, linhas):
    celulas = [[c.strip() for c in l.strip().strip('|').split('|')] for l in linhas
               if not re.fullmatch(r'\|?[\s:|-]+\|?', l.strip())]
    t = doc.add_table(rows=len(celulas), cols=max(len(c) for c in celulas))
    t.style = 'Table Grid'
    for i, linha in enumerate(celulas):
        for j, valor in enumerate(linha):
            p = t.cell(i, j).paragraphs[0]
            texto_rico(p, f'**{valor}**' if i == 0 and valor else valor)


def converte(md, docx):
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    linhas = Path(md).read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(linhas):
        l = linhas[i]
        if not l.strip():
            i += 1
        elif l.startswith('#'):
            nivel = len(l) - len(l.lstrip('#'))
            doc.add_heading(FONTE.sub('', l.lstrip('#')).strip(), 0 if nivel == 1 else nivel - 1)
            i += 1
        elif l.lstrip().startswith('|'):
            bloco = []
            while i < len(linhas) and linhas[i].lstrip().startswith('|'):
                bloco.append(linhas[i])
                i += 1
            tabela(doc, bloco)
        elif ITEM.match(l):
            recuo, marca, texto = ITEM.match(l).groups()
            i += 1
            while i < len(linhas) and linhas[i].strip() and not ITEM.match(linhas[i]) \
                    and not linhas[i].lstrip().startswith(('|', '#')) and len(linhas[i]) - len(linhas[i].lstrip()) > len(recuo):
                texto += ' ' + linhas[i].strip()
                i += 1
            nivel = min(len(recuo) // 2, 2)
            if marca[0].isdigit():
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18 + 18 * nivel)
                p.paragraph_format.first_line_indent = Pt(-14)
                p.add_run(marca + ' ')
            else:
                p = doc.add_paragraph(style='List Bullet' + (f' {nivel + 1}' if nivel else ''))
            texto_rico(p, texto)
        else:
            texto = l.strip()
            i += 1
            while i < len(linhas) and linhas[i].strip() and not ITEM.match(linhas[i]) \
                    and not linhas[i].lstrip().startswith(('|', '#')):
                texto += ' ' + linhas[i].strip()
                i += 1
            texto_rico(doc.add_paragraph(), texto)
    doc.save(docx)


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    origem = Path(sys.argv[1])
    converte(origem, Path(sys.argv[2]) if len(sys.argv) > 2 else origem.with_suffix('.docx'))
