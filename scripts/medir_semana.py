"""Linha semanal de uso do Claude Code a partir das transcrições (protótipo da R1, 06/10; completo no F2b, 09/10: a planilha, o % do limite semanal e o gancho de abertura que roda esta medição). Só números."""
# Lê ~/.claude/projects/**/*.jsonl (principais e subagentes). O texto das mensagens é só testado (começa com tal
# marcador?) e comparado por HASH (sha1, nunca impresso); nada de texto, caminho, comando ou título é impresso nem
# gravado. Saída: contagens, durações, datas, IDs de modelo, níveis de esforço, nomes de pasta de projeto e de projeto.
#
#     python medir_semana.py                              # semana 2026-09-29 .. 2026-10-05 (fuso local)
#     python medir_semana.py --inicio 2026-10-06 --dias 7
#     python medir_semana.py --csv uso-semanal.csv --ctx0-csv ctx0-por-projeto.csv --mapa pastas-projetos.csv
#         (a linha da semana troca a de mesmo semana_inicio e pc, ou entra no fim; o ctx0 troca as da semana)
#     python medir_semana.py --planilha ../metricas/uso.xlsx --csv ../metricas/uso-semanal.csv
#         (só refaz a planilha a partir do CSV e do metricas/metas.json; não lê transcrição)
#
# Quem roda sozinho: o gancho de abertura do nucleo (plugins/nucleo/ganchos/semana.py), em segundo plano, quando falta
# a linha de uma semana já fechada deste PC; ele refaz a planilha em seguida. O pct_limite_semanal vem do arquivo que a
# statusline grava (~/.claude/kit-local/limites.jsonl): o maior % semanal visto dentro da semana.
#
# Regras (pesquisas/2026-10-06_medicao-semanal-de-uso.md, seção 4):
#   humano estrito  = H (digitada) + HQ (digitada com o Claude ocupado, entra como attachment queued_command) + HC (comando)
#   presença ampla  = estrito + HA (resposta a AskUserQuestion/ExitPlanMode, ou recusa de permissão) + HI (interrupção)
#   trabalho do Claude (W) = chamada do modelo (assistant, não sintética) + tool_result de ferramenta; principal e subagentes
#   trecho C      = corrida de W de uma sessão (lacuna <= 15 min), que começa na presença ampla e é cortada por ela
#   trecho global = a mesma corrida, com os W e a presença ampla de TODAS as sessões numa só linha do tempo
#   limite batido = mensagem sintética de erro da API com quotaLimits (arquivos principais e de subagentes)
#   esforço       = campo `effort` (topo) da entrada assistant; chamada = message.id único (principal + subagente)
import os, re, sys, csv, json, time, bisect, hashlib, argparse, statistics, collections
import datetime as dt

BASE = os.path.join(os.path.expanduser('~'), '.claude', 'projects')
LIMITES = os.path.join(os.path.expanduser('~'), '.claude', 'kit-local', 'limites.jsonl')
KIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RX_INT = re.compile(r'^-?\d+$')
VERSAO = 'r1-1'
LIM_RETORNO = 30 * 60      # retorno: humano volta depois de > 30 min sem evento
LIM_SENTADA = 20 * 60      # sentada: mensagens do humano com intervalo <= 20 min
JANELA_CONC = 10 * 60      # janela da concorrencia
GAP_CONTINUO = 15 * 60     # trecho continuo: lacuna maxima entre eventos de trabalho
LIM_RETORNO_VAR = (20, 60)  # minutos: variantes do retorno (colunas retornos_20min_estrito e retornos_60min_estrito)
LIM_SENTADA_VAR = (10, 30)  # minutos: variantes da sentada (colunas sentadas_10min_* e sentadas_30min_*)
RX_ENUM = re.compile(r'^[A-Za-z][A-Za-z0-9_\-\.]{0,40}$')
RX_MODELO = re.compile(r'^[A-Za-z0-9_.\-\[\]<>]{1,60}$')
PRES_ESTRITA = ('H', 'HQ', 'HC')
PRES_AMPLA = ('H', 'HQ', 'HC', 'HA', 'HI')
W = ('A', 'T')
# ID do modelo -> coluna de chamadas (principal + subagente)
MODELOS = (('claude-opus-5-5', 'chamadas_opus55'), ('claude-opus-5', 'chamadas_opus5'),
           ('claude-sonnet-5-5', 'chamadas_sonnet55'), ('claude-sonnet-5', 'chamadas_sonnet5'),
           ('claude-fable-5-1', 'chamadas_fable51'), ('claude-haiku-4-5-20251001', 'chamadas_haiku45'))
CABECALHO = ['semana_inicio', 'sessoes_com_chamada', 'chamadas', 'chamadas_subagente', 'tok_entrada', 'tok_cache_lido',
             'tok_cache_gravado', 'tok_saida', 'ctx_mediano_principal', 'ctx0_mediano', 'msgs_humano_estrito',
             'respostas_perguntas', 'retornos_30min_estrito', 'retornos_30min_amplo', 'maior_trecho_A_min',
             'maior_trecho_B_min', 'maior_trecho_C_min', 'conc_max_humano', 'conc_max_bruta',
             'horas_2mais_humano_janela10', 'sentadas_estrito', 'sentadas_min_estrito'] + [c for _, c in MODELOS] + [
            'maior_trecho_global_min', 'trecho_global_mediana_min', 'retornos_20min_estrito', 'retornos_60min_estrito',
            'sentadas_10min_blocos', 'sentadas_10min_min', 'sentadas_30min_blocos', 'sentadas_30min_min',
            'chamadas_esforco_max', 'chamadas_esforco_xhigh', 'limites_batidos', 'pct_limite_semanal', 'pc', 'versao']
CABECALHO_CTX0 = ['semana_inicio', 'projeto', 'sessoes', 'ctx0_mediana', 'ctx0_min', 'ctx0_max']


def ts_de(s):
    return dt.datetime.fromisoformat(s.replace('Z', '+00:00')).timestamp()   # 'Z' so e aceito no Python 3.11+


def texto(c):
    if isinstance(c, str):
        return c
    if isinstance(c, list):
        for b in c:
            if isinstance(b, dict) and b.get('type') == 'text':
                return b.get('text') or ''
    return ''


def hsh(s):
    return hashlib.sha1((s or '').strip().encode('utf-8', 'replace')).hexdigest()


def enum(v):
    if v is None:
        return 'null'
    return v if isinstance(v, str) and RX_ENUM.match(v) else '<fora-da-regra>'


def modelo(v):
    return v if isinstance(v, str) and RX_MODELO.match(v) else '<fora-da-regra>'


def pct(v, p):
    v = sorted(v)
    return v[int(p * (len(v) - 1))] if v else 0


def med(v):
    return statistics.median(v) if v else 0


def dia_local(t):
    return dt.datetime.fromtimestamp(t).strftime('%Y-%m-%d')


def hm(seg):
    seg = int(round(seg))
    return '%dh%02dmin' % (seg // 3600, (seg % 3600) // 60)


def corridas(seq, gap, t_ini, t_fim):
    """seq = [(t, 'W' ou 'P')] em ordem (W = trabalho do Claude, P = presenca ampla do humano). Devolve
    [(duracao_s, inicio, 0)] das corridas de W com lacuna <= gap: a corrida comeca no P ou no 1o W depois de uma
    lacuna maior, e qualquer P a corta. Conta so a de duracao > 0 cujo inicio cai em [t_ini, t_fim)."""
    saida = []
    ini = ult = None
    for t, k in seq:
        if k == 'P':
            if ini is not None and ult > ini and t_ini <= ini < t_fim:
                saida.append((ult - ini, ini, 0))
            ini = ult = t
        else:
            if ini is None or t - ult > gap:
                if ini is not None and ult > ini and t_ini <= ini < t_fim:
                    saida.append((ult - ini, ini, 0))
                ini = ult = t
            else:
                ult = t
    if ini is not None and ult > ini and t_ini <= ini < t_fim:
        saida.append((ult - ini, ini, 0))
    return saida


def pct_semanal(caminho, t_ini, t_fim):
    """O maior % do limite semanal (rate_limits.seven_day) que a statusline anotou dentro da semana; '' sem dado."""
    vals = []
    for f in (caminho + '.1', caminho):
        if not os.path.exists(f):
            continue
        with open(f, encoding='utf-8') as fh:
            for ln in fh:
                try:
                    r = json.loads(ln)
                except ValueError:
                    continue
                if isinstance(r.get('sem'), (int, float)) and t_ini <= float(r.get('ts', 0)) < t_fim:
                    vals.append(round(r['sem']))
    return max(vals) if vals else ''


def _num(v):
    return int(v) if RX_INT.match(v or '') else None


def gerar_planilha(xlsx, caminho_csv, caminho_metas):
    """metricas/uso.xlsx: a aba "semanas" (o CSV) e a aba "regua" (um gráfico de linha por métrica, com a meta).

    Um gráfico por métrica, nunca dois eixos; uma cor por PC e a meta em cinza tracejado (skill dataviz). Fica fora do
    git (metricas/*.xlsx): cada PC a refaz."""
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font
    with open(caminho_csv, encoding='utf-8', newline='') as fh:
        linhas = [r for r in csv.reader(fh) if r]
    with open(caminho_metas, encoding='utf-8') as fh:
        metas = json.load(fh)
    cab, dados = linhas[0], linhas[1:]
    i_pc = cab.index('pc')
    wb = Workbook()
    ws = wb.active
    ws.title = 'semanas'
    ws.append(cab)
    for r in dados:
        ws.append([_num(x) if _num(x) is not None else x for x in r])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = 'B2'
    rg = wb.create_sheet('regua')
    rg['A1'] = 'A régua (metricas/metas.json): o valor de cada semana, por PC, e a meta. Base: a semana de %s.' % (
        metas.get('base_semana', ''))
    rg['A1'].font = Font(bold=True)
    pcs = sorted({r[i_pc] for r in dados})
    semanas = sorted({r[0] for r in dados})
    cores = ['2A78D6', '1BAF7A', 'E6A100']          # series-1..3 da paleta de referência, sempre nesta ordem
    lin = 3
    for m in metas.get('regua', []):
        if m['coluna'] not in cab:
            continue
        j = cab.index(m['coluna'])
        tem_meta = m.get('meta') is not None
        rg.cell(lin, 1, m['titulo']).font = Font(bold=True)
        rg.cell(lin + 1, 1, 'semana')
        for k, p in enumerate(pcs):
            rg.cell(lin + 1, 2 + k, p)
        if tem_meta:
            rg.cell(lin + 1, 2 + len(pcs), 'meta (%s)' % m.get('regra', ''))
        for s_i, s in enumerate(semanas):
            rg.cell(lin + 2 + s_i, 1, s)
            for k, p in enumerate(pcs):
                v = next((r[j] for r in dados if r[0] == s and r[i_pc] == p), '')
                rg.cell(lin + 2 + s_i, 2 + k, _num(v))
            if tem_meta:
                rg.cell(lin + 2 + s_i, 2 + len(pcs), m['meta'])
        ult = lin + 1 + len(semanas)
        ch = LineChart()
        ch.title = m['titulo'].split(' (')[0]          # o nome completo fica na célula acima da tabela
        ch.height, ch.width = 6.5, 14
        ch.legend.position = 'b'
        ch.add_data(Reference(rg, min_col=2, max_col=1 + len(pcs) + tem_meta, min_row=lin + 1, max_row=ult),
                    titles_from_data=True)
        ch.set_categories(Reference(rg, min_col=1, min_row=lin + 2, max_row=ult))
        for k, s in enumerate(ch.series):
            e_meta = tem_meta and k == len(pcs)
            cor = '52514E' if e_meta else cores[k % len(cores)]
            s.smooth = False
            s.graphicalProperties.line.width = 19050 if e_meta else 25400       # 1,5 pt e 2 pt
            s.graphicalProperties.line.solidFill = cor
            if e_meta:
                s.graphicalProperties.line.dashStyle = 'dash'
                s.marker.symbol = 'none'
            else:
                s.marker.symbol = 'circle'
                s.marker.size = 7
                s.marker.graphicalProperties.solidFill = cor
                s.marker.graphicalProperties.line.solidFill = cor
        rg.add_chart(ch, 'G%d' % lin)
        lin = max(ult + 2, lin + 15)
    rg.column_dimensions['A'].width = 14
    os.makedirs(os.path.dirname(os.path.abspath(xlsx)), exist_ok=True)
    tmp = xlsx + '.tmp.xlsx'
    wb.save(tmp)
    os.replace(tmp, xlsx)
    print('planilha: %d semana(s), %d PC(s), %d gráfico(s) -> %s' % (len(semanas), len(pcs), len(rg._charts), xlsx))
    return 0


def carregar_mapa(caminho):
    """CSV pasta,projeto -> dict {pasta em minusculas: projeto} (o Windows nao diferencia maiusculas)."""
    mapa = {}
    with open(caminho, encoding='utf-8', newline='') as fh:
        for r in csv.DictReader(fh):
            p = (r.get('pasta') or '').strip()
            j = (r.get('projeto') or '').strip()
            if p and j:
                mapa[p.lower()] = j
    return mapa


def projeto_de(pasta, mapa):
    """Sem mapa: o proprio nome da pasta. Com mapa: o projeto, ou 'desconhecido' se a pasta nao esta nele."""
    if mapa is None:
        return pasta
    return mapa.get(pasta.lower(), 'desconhecido')


def gravar_csv(caminho, cabecalho, novas, trocar, ordem, manter=()):
    """CSV (UTF-8, virgula, fim de linha LF): mantem as linhas antigas, tira as que trocar(linha) marca, poe as novas
    e ordena pelas colunas `ordem`. Nas colunas `manter`, a celula vazia da linha nova herda a da linha que ela
    substitui (valor digitado depois, como o pct_limite_semanal). Cabecalho diferente do esperado: nao grava e devolve
    False. Escreve num arquivo temporario e troca (o CSV e o historico quando a varredura de 30 dias apaga as transcricoes)."""
    antigas, trocadas = [], []
    if os.path.exists(caminho):
        with open(caminho, encoding='utf-8', newline='') as fh:
            todas = [r for r in csv.reader(fh) if r]
        if todas:
            if todas[0] != cabecalho:
                return False
            for r in todas[1:]:
                (trocadas if trocar(r) else antigas).append(r)
    novas = [list(r) for r in novas]
    if trocadas:
        for r in novas:
            for c in manter:
                i = cabecalho.index(c)
                if r[i] == '' and len(trocadas[0]) > i:
                    r[i] = trocadas[0][i]
    linhas = sorted(antigas + novas, key=lambda r: tuple(r[i] for i in ordem))
    os.makedirs(os.path.dirname(os.path.abspath(caminho)), exist_ok=True)
    tmp = caminho + '.tmp'
    with open(tmp, 'w', encoding='utf-8', newline='') as fh:
        w = csv.writer(fh, lineterminator='\n')
        w.writerow(cabecalho)
        w.writerows(linhas)
    os.replace(tmp, caminho)
    return True


# ----------------------------------------------------------------------------- leitura de um arquivo
def ler_arquivo(caminho, eh_sub, sessao, projeto, t_ini, t_fim, ctrl):
    """Devolve eventos [(ts, classe)], chamadas {msg_id: dict}, enfileirados [(ts, hash)], humanos com hash."""
    eventos, chamadas, enq, hum = [], {}, [], []
    tool_nome = {}
    for linha in open(caminho, encoding='utf-8', errors='replace'):
        try:
            d = json.loads(linha)
        except Exception:
            ctrl['erro_json'] += 1
            continue
        t = d.get('type')
        tsr = d.get('timestamp')
        if not isinstance(tsr, str):
            continue
        try:
            ts = ts_de(tsr)
        except Exception:
            continue
        if t == 'assistant':
            m = d.get('message') or {}
            mdl = m.get('model')
            c = m.get('content')
            if isinstance(c, list):
                for b in c:
                    if isinstance(b, dict) and b.get('type') == 'tool_use' and b.get('id'):
                        tool_nome[b['id']] = b.get('name') if b.get('name') in ('AskUserQuestion', 'ExitPlanMode') else 'x'
            if not mdl or mdl == '<synthetic>':
                eventos.append((ts, 'E'))
                if d.get('isApiErrorMessage'):
                    ctrl['erro_api'][(dia_local(ts), 'com_quota' if d.get('quotaLimits') else 'sem_quota', 'sub' if eh_sub else 'principal')] += 1
                    if d.get('quotaLimits'):
                        ctrl['limites'].append((ts, eh_sub))      # limite batido: erro sintetico da API com quotaLimits
                continue
            eventos.append((ts, 'A'))
            u = m.get('usage') or {}
            chamadas[m.get('id') or d.get('uuid')] = dict(
                ts=ts, modelo=modelo(mdl), esforco=enum(d.get('effort')), sub=eh_sub, sessao=sessao, projeto=projeto,
                ent=u.get('input_tokens') or 0, lido=u.get('cache_read_input_tokens') or 0,
                grav=u.get('cache_creation_input_tokens') or 0, saida=u.get('output_tokens') or 0)
        elif t == 'user':
            m = d.get('message') or {}
            c = m.get('content')
            trs = [b for b in c if isinstance(b, dict) and b.get('type') == 'tool_result'] if isinstance(c, list) else []
            if trs:
                pergunta = any(tool_nome.get(b.get('tool_use_id'), 'x') in ('AskUserQuestion', 'ExitPlanMode') for b in trs)
                if eh_sub:
                    cls = 'T'
                elif pergunta or d.get('toolDenialKind') == 'user-rejected':
                    cls = 'HA'
                else:
                    cls = 'T'
                eventos.append((ts, cls))
                continue
            if eh_sub:
                eventos.append((ts, 'X'))
                continue
            if d.get('isCompactSummary') or d.get('isMeta'):
                eventos.append((ts, 'X'))
                continue
            txt = texto(c)
            o = d.get('origin')
            kind = o.get('kind') if isinstance(o, dict) else None
            if kind == 'task-notification' or txt.startswith('<task-notification'):
                cls = 'N'
            elif kind == 'peer':
                cls = 'X'
            elif txt.startswith('[Request interrupted'):
                cls = 'HI'
            elif txt.startswith('<command-name>') or txt.startswith('<command-message>'):
                cls = 'HC'
            elif txt.startswith('<local-command-'):
                cls = 'X'
            elif kind == 'human':
                cls = 'H'
                hum.append((ts, hsh(txt)))
            elif kind is None and d.get('promptSource') == 'sdk':
                cls = 'HL'      # legado: versao antiga, sem campo origin; nao entra no conjunto do humano
            else:
                cls = 'U'
            eventos.append((ts, cls))
        elif t == 'attachment':
            if eh_sub:
                eventos.append((ts, 'X'))
                continue
            at = d.get('attachment') or {}
            if at.get('type') == 'queued_command':
                o = at.get('origin')
                kind = o.get('kind') if isinstance(o, dict) else None
                if kind == 'human':
                    eventos.append((ts, 'HQ'))
                    hum.append((ts, hsh(texto(at.get('prompt')))))
                elif kind == 'task-notification' or at.get('commandMode') == 'task-notification':
                    eventos.append((ts, 'N'))
                else:
                    eventos.append((ts, 'X'))
            else:
                eventos.append((ts, 'X'))
        elif t == 'queue-operation':
            eventos.append((ts, 'X'))
            if d.get('operation') == 'enqueue' and isinstance(d.get('content'), str):
                enq.append((ts, hsh(d['content'])))
        elif t in ('system', 'file-history-delta', 'frame-link'):
            eventos.append((ts, 'X'))
    return eventos, chamadas, enq, hum


# ----------------------------------------------------------------------------- principal
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--inicio', default='2026-09-29')
    ap.add_argument('--dias', type=int, default=7)
    ap.add_argument('--csv', default='')
    ap.add_argument('--retorno-min', type=float, default=30)
    ap.add_argument('--sentada-min', type=float, default=20)
    ap.add_argument('--gap-continuo-min', type=float, default=15)
    ap.add_argument('--so-linha', action='store_true', help='imprime so a linha semanal (cabecalho + valores)')
    ap.add_argument('--pc', default='notebook', help='rotulo do PC na coluna pc da linha semanal')
    ap.add_argument('--ctx0-csv', default='', help='CSV longo do ctx0 por projeto (troca as linhas da semana)')
    ap.add_argument('--mapa', default='', help='CSV pasta,projeto (pastas-projetos.csv); sem ele, o ctx0 sai por pasta')
    ap.add_argument('--limites', default=LIMITES, help='o arquivo da statusline com o %% do limite semanal')
    ap.add_argument('--planilha', default='', help='so refaz a planilha (xlsx) a partir do --csv e do --metas')
    ap.add_argument('--metas', default=os.path.join(KIT, 'metricas', 'metas.json'))
    a = ap.parse_args()
    if a.planilha:
        return gerar_planilha(a.planilha, a.csv or os.path.join(KIT, 'metricas', 'uso-semanal.csv'), a.metas)
    if not RX_ENUM.match(a.pc):
        ap.error('--pc: so letras, digitos, _, - e . (rotulo curto)')
    global LIM_RETORNO, LIM_SENTADA, GAP_CONTINUO
    LIM_RETORNO = a.retorno_min * 60
    LIM_SENTADA = a.sentada_min * 60
    GAP_CONTINUO = a.gap_continuo_min * 60
    mapa = carregar_mapa(a.mapa) if a.mapa else None
    t_prog = time.perf_counter()
    ano, mes, dd = (int(x) for x in a.inicio.split('-'))
    d0 = dt.datetime(ano, mes, dd)
    t_ini = d0.timestamp()
    t_fim = (d0 + dt.timedelta(days=a.dias)).timestamp()
    dias = [(d0 + dt.timedelta(days=i)).strftime('%Y-%m-%d') for i in range(a.dias)]

    # arquivos candidatos: mtime >= inicio (um arquivo nao alterado antes do inicio nao tem evento na semana)
    principais, subs = [], []
    ign = 0
    mb_lidos = [0]
    for pasta in sorted(os.listdir(BASE)):
        pp = os.path.join(BASE, pasta)
        if not os.path.isdir(pp):
            continue
        for raiz, _, arqs in os.walk(pp):
            partes = os.path.relpath(raiz, pp).split(os.sep)
            for x in arqs:
                if not x.endswith('.jsonl'):
                    continue
                f = os.path.join(raiz, x)
                if os.path.getmtime(f) < t_ini:
                    ign += 1
                    continue
                mb_lidos[0] += os.path.getsize(f)
                if 'subagents' in partes:
                    sess = partes[partes.index('subagents') - 1]
                    subs.append((f, sess, pasta))
                elif partes == ['.']:
                    principais.append((f, os.path.splitext(x)[0], pasta))
    ctrl = dict(erro_json=0, erro_api=collections.Counter(), limites=[])
    todas_chamadas = {}
    dup = [0, 0]  # [id repetido em outro arquivo, total]
    por_sessao = collections.defaultdict(lambda: dict(eventos=[], humanos=[], projeto=None, primeira=None))
    enq_por_sessao = {}
    t_leitura = time.perf_counter()
    for lista, eh_sub in ((principais, False), (subs, True)):   # principais primeiro: o id fica com o principal
        for f, sess, proj in lista:
            ev, ch, enq, hum = ler_arquivo(f, eh_sub, sess, proj, t_ini, t_fim, ctrl)
            s = por_sessao[sess]
            s['projeto'] = s['projeto'] or proj
            s['eventos'].extend(ev)
            if not eh_sub:
                s['humanos'].extend(hum)
                enq_por_sessao[sess] = enq
            for mid, c in ch.items():
                dup[1] += 1
                if mid in todas_chamadas:
                    dup[0] += 1
                    continue
                todas_chamadas[mid] = c
    t_leitura = time.perf_counter() - t_leitura

    # ---------------- ajuste do instante digitado: mensagem digitada com o Claude ocupado tem 'enqueue' antes
    ajuste = []
    humanos_ev = {}  # sess -> lista de (ts_ajustado, classe) para H/HQ
    for sess, s in por_sessao.items():
        enq = collections.defaultdict(list)
        for tt, hh in enq_por_sessao.get(sess, []):
            enq[hh].append(tt)
        ev2 = []
        # reclassifica: a lista 'humanos' tem (ts, hash) na ordem de leitura; casamos por ts com os eventos H/HQ
        hs = {round(tt, 3): hh for tt, hh in s['humanos']}
        novos = []
        for ts, cls in s['eventos']:
            if cls in ('H', 'HQ'):
                hh = hs.get(round(ts, 3))
                cand = [x for x in enq.get(hh, []) if x <= ts and ts - x <= 7200] if hh else []
                if cand:
                    t2 = max(cand)
                    ajuste.append(ts - t2)
                    ts = t2
            novos.append((ts, cls))
        s['eventos'] = sorted(novos)
    ev_total = sum(len(s['eventos']) for s in por_sessao.values())
    # sessao do humano = o arquivo tem ao menos 1 mensagem com origin.kind == 'human' (digitada ou enfileirada)
    sess_humana = {k for k, s in por_sessao.items() if any(c in ('H', 'HQ') for t, c in s['eventos'])}
    # duracao de trabalho de cada sessao (1o a ultimo evento de trabalho, arquivo todo, com subagentes)
    dur_trab = {}
    for k, s in por_sessao.items():
        w = [t for t, c in s['eventos'] if c in W]
        dur_trab[k] = (max(w) - min(w)) if w else 0

    # ---------------- CLASSES (controle)
    cont_cls = collections.Counter()
    for s in por_sessao.values():
        for ts, cls in s['eventos']:
            if t_ini <= ts < t_fim:
                cont_cls[cls] += 1

    # ---------------- 1) TOKENS / CHAMADAS na semana (por sessao, por modelo)
    ch_sem = [c for c in todas_chamadas.values() if t_ini <= c['ts'] < t_fim]
    tok = collections.defaultdict(lambda: collections.Counter())
    for c in ch_sem:
        for lado in ('total', 'sub' if c['sub'] else 'principal'):
            tok[lado]['chamadas'] += 1
            tok[lado]['entrada'] += c['ent']
            tok[lado]['cache_lido'] += c['lido']
            tok[lado]['cache_gravado'] += c['grav']
            tok[lado]['saida'] += c['saida']
    por_modelo = collections.Counter()
    por_modelo_esf = collections.Counter()
    por_modelo_sub = collections.Counter()
    for c in ch_sem:
        (por_modelo_sub if c['sub'] else por_modelo)[c['modelo']] += 1
        por_modelo_esf[(c['modelo'], c['esforco'])] += 1
    sess_tok = collections.defaultdict(lambda: collections.Counter())
    for c in ch_sem:
        k = c['sessao']
        sess_tok[k]['chamadas'] += 1
        sess_tok[k]['entrada'] += c['ent']
        sess_tok[k]['lido'] += c['lido']
        sess_tok[k]['grav'] += c['grav']
        sess_tok[k]['saida'] += c['saida']
        sess_tok[k]['sub'] += 1 if c['sub'] else 0

    # ---------------- 2) CONTEXTO mediano por chamada (principais) e ctx0 por projeto
    ctx_pr = [c['ent'] + c['lido'] + c['grav'] for c in ch_sem if not c['sub']]
    ctx_sub = [c['ent'] + c['lido'] + c['grav'] for c in ch_sem if c['sub']]
    # ctx0: 1a chamada (principal) da sessao, so sessoes cuja 1a chamada cai na semana
    prim = {}
    for c in todas_chamadas.values():
        if c['sub']:
            continue
        k = c['sessao']
        if k not in prim or c['ts'] < prim[k]['ts']:
            prim[k] = c
    ctx0_proj = collections.defaultdict(list)     # so sessoes do humano (tem mensagem com origin.kind == human)
    ctx0_todas = []                                # todas as sessoes principais, sem filtro
    for k, c in prim.items():
        if t_ini <= c['ts'] < t_fim:
            ctx0_todas.append(c['ent'] + c['lido'] + c['grav'])
            if k in sess_humana:
                ctx0_proj[c['projeto']].append(c['ent'] + c['lido'] + c['grav'])
    # ctx0 dos subagentes (1a chamada de cada arquivo de subagente) nao e separavel aqui sem id de arquivo; omitido

    # ---------------- 3) HUMANOS por sessao (na semana)
    hum_sess = {}
    for sess, s in por_sessao.items():
        n = collections.Counter(cls for ts, cls in s['eventos'] if cls in PRES_AMPLA and t_ini <= ts < t_fim)
        if n:
            hum_sess[sess] = n
    msgs_estrito = [sum(n[c] for c in PRES_ESTRITA) for n in hum_sess.values()]

    # ---------------- 4) RETORNOS (humano volta apos > 30 min sem evento) e 5) TRECHOS
    gaps_h = []          # lacuna antes de cada evento humano (exceto o 1o da sessao), variante 2
    retornos = collections.Counter()
    retornos_var = collections.Counter()     # (conjunto, limiar em min) -> n, so variante 2 (LIM_RETORNO_VAR)
    retornos_dia = collections.Counter()
    inicios = 0
    trechos_A, trechos_B, trechos_C = [], [], []
    seq_global = []      # (t, 'W' ou 'P') de TODAS as sessoes, para o trecho global
    for sess, s in por_sessao.items():
        ev = s['eventos']
        if not ev:
            continue
        ts_all = [t for t, c in ev]
        ts_v2 = [t for t, c in ev if c in ('A', 'T', 'N', 'E') or c in PRES_AMPLA]
        ts_W = [t for t, c in ev if c in W]
        pres_e = [t for t, c in ev if c in PRES_ESTRITA]
        pres_a = [t for t, c in ev if c in PRES_AMPLA]
        primeira_pres = pres_a[0] if pres_a else None
        for conj, nome_conj in ((PRES_ESTRITA, 'estrito'), (PRES_AMPLA, 'amplo')):
            pres = [t for t, c in ev if c in conj]
            for h in pres:
                if not (t_ini <= h < t_fim):
                    continue
                if h == pres[0]:
                    if nome_conj == 'estrito':
                        inicios += 1
                    continue
                for nome_v, lst in (('v1_qualquer_entrada', ts_all), ('v2_so_trabalho_e_humano', ts_v2)):
                    i = bisect.bisect_left(lst, h) - 1
                    if i < 0:
                        continue
                    gap = h - lst[i]
                    if nome_v == 'v2_so_trabalho_e_humano' and nome_conj == 'amplo':
                        gaps_h.append(gap)
                    if gap > LIM_RETORNO:
                        retornos[(nome_conj, nome_v)] += 1
                        if nome_v == 'v2_so_trabalho_e_humano':
                            retornos_dia[(nome_conj, dia_local(h))] += 1
                    if nome_v == 'v2_so_trabalho_e_humano':
                        for lim in LIM_RETORNO_VAR:
                            if gap > lim * 60:
                                retornos_var[(nome_conj, lim)] += 1
        # trechos A (fronteira = humano estrito) e B (fronteira = presenca ampla)
        for conj, saida in ((pres_e, trechos_A), (pres_a, trechos_B)):
            for k, h in enumerate(conj):
                if not (t_ini <= h < t_fim):
                    continue
                nxt = conj[k + 1] if k + 1 < len(conj) else float('inf')
                j0 = bisect.bisect_right(ts_W, h)
                j1 = bisect.bisect_left(ts_W, nxt)
                if j1 > j0:
                    saida.append((ts_W[j1 - 1] - h, h, j1 - j0))
        # trechos C: corridas continuas (lacuna <= GAP_CONTINUO), cortadas por evento humano (presenca ampla)
        seq = sorted([(t, 'W') for t in ts_W] + [(t, 'P') for t in pres_a])
        trechos_C.extend(corridas(seq, GAP_CONTINUO, t_ini, t_fim))
        seq_global.extend(seq)
    # trecho global: os W e a presenca ampla de TODAS as sessoes numa so linha do tempo (a mesma regra do C)
    trechos_G = corridas(sorted(seq_global), GAP_CONTINUO, t_ini, t_fim)

    # ---------------- 6) CONCORRENCIA (sessoes com evento de trabalho na mesma janela de 10 min)
    def concorrencia(filtro):
        bucket = collections.defaultdict(set)   # (dia, indice10min) -> sessoes
        hora = collections.defaultdict(set)     # (dia, hora) -> sessoes
        for sess, s in por_sessao.items():
            if not filtro(sess):
                continue
            for ts, cls in s['eventos']:
                if cls in W and t_ini <= ts < t_fim:
                    x = dt.datetime.fromtimestamp(ts)
                    dia = x.strftime('%Y-%m-%d')
                    bucket[(dia, x.hour * 6 + x.minute // 10)].add(sess)
                    hora[(dia, x.hour)].add(sess)
        conc = {d: 0 for d in dias}
        for (d, b), ss in bucket.items():
            conc[d] = max(conc[d], len(ss))
        solto = collections.Counter(d for (d, h), ss in hora.items() if len(ss) >= 2)
        horas_e = {(d, b // 6) for (d, b), ss in bucket.items() if len(ss) >= 2}
        estrito = collections.Counter(d for d, h in horas_e)
        return conc, solto, estrito, len(hora)
    c_bruta = concorrencia(lambda k: True)
    c_hum = concorrencia(lambda k: k in sess_humana)
    c_hum10 = concorrencia(lambda k: k in sess_humana and dur_trab.get(k, 0) >= 600)
    conc_dia, h2_solto, h2_estrito, horas_com_trabalho = c_hum
    # pico por 'span' (como o session-report): sessoes com [primeiro,ultimo] evento cobrindo a janela
    pico_span = {d: 0 for d in dias}
    spans = {}
    for sess, s in por_sessao.items():
        ts_W = [t for t, c in s['eventos'] if c in W]
        if ts_W:
            spans[sess] = (min(ts_W), max(ts_W))
    for d in dias:
        base = dt.datetime.strptime(d, '%Y-%m-%d').timestamp()
        b = [0] * 144
        for sess, (a0, a1) in spans.items():
            if a1 < base or a0 >= base + 86400:
                continue
            lo = max(0, int((max(a0, base) - base) // 600))
            hi = min(144, int((min(a1, base + 86399) - base) // 600) + 1)
            for i in range(lo, hi):
                b[i] += 1
        pico_span[d] = max(b)

    # ---------------- 7) SENTADAS
    def sentadas(conj, lim):
        tt = sorted(ts for s in por_sessao.values() for ts, c in s['eventos'] if c in conj and t_ini <= ts < t_fim)
        blocos = []
        for t in tt:
            if blocos and t - blocos[-1][1] <= lim:
                blocos[-1][1] = t
                blocos[-1][2] += 1
            else:
                blocos.append([t, t, 1])
        return blocos
    sent_e = sentadas(PRES_ESTRITA, LIM_SENTADA)
    sent_a = sentadas(PRES_AMPLA, LIM_SENTADA)
    sent_var = {m: sentadas(PRES_ESTRITA, m * 60) for m in LIM_SENTADA_VAR}    # variantes (so estrito)

    # ---------------- 8) ESFORCO e LIMITES BATIDOS (na semana; principal + subagente)
    esf_tot = collections.Counter(c['esforco'] for c in ch_sem)
    lim_sem = [sub for ts, sub in ctrl['limites'] if t_ini <= ts < t_fim]

    # ---------------- SAIDA
    def out(*x):
        if not a.so_linha:
            print(*x)
    dur_prog = time.perf_counter() - t_prog
    out('SEMANA %s .. %s (fuso local UTC-3; %d dias)' % (dias[0], dias[-1], a.dias))
    out('arquivos lidos: principais %d, subagentes %d (%.1f MB); ignorados por mtime anterior ao inicio: %d; erros de JSON: %d' % (
        len(principais), len(subs), mb_lidos[0] / 1e6, ign, ctrl['erro_json']))
    out('eventos com timestamp lidos: %d; message.id repetido entre arquivos (descartado): %d de %d' % (ev_total, dup[0], dup[1]))
    out('classes de evento na semana (controle):', dict(sorted(cont_cls.items())))
    out('ajuste do instante digitado pelo enqueue: n=%d  mediana=%.1fs p90=%.1fs max=%.1fs' % (
        len(ajuste), med(ajuste), pct(ajuste, .9), max(ajuste) if ajuste else 0))

    out('\n[1] TOKENS E CHAMADAS NA SEMANA (chamada = message.id unico; uso do ultimo bloco)')
    for lado in ('total', 'principal', 'sub'):
        x = tok[lado]
        tt = x['entrada'] + x['cache_lido'] + x['cache_gravado'] + x['saida']
        out('  %-9s chamadas %6d | entrada %12d | cache lido %14d | cache gravado %12d | saida %10d | total %14d' % (
            lado, x['chamadas'], x['entrada'], x['cache_lido'], x['cache_gravado'], x['saida'], tt))
    tt = tok['total']
    out('  cache lido / (entrada+lido+gravado) = %.1f%%' % (100.0 * tt['cache_lido'] / max(1, tt['entrada'] + tt['cache_lido'] + tt['cache_gravado'])))
    out('  sessoes com chamada na semana: %d (so principais: %d)' % (len(sess_tok), len([k for k in sess_tok if sess_tok[k]['chamadas'] > sess_tok[k]['sub']])))
    soma_s = sorted(((v['entrada'] + v['lido'] + v['grav'] + v['saida']) for v in sess_tok.values()))
    out('  total de tokens por sessao: mediana %d  p90 %d  max %d' % (med(soma_s), pct(soma_s, .9), soma_s[-1] if soma_s else 0))
    for nome, cond in (('sessoes do humano (tem mensagem origin.kind=human)', lambda k: k in sess_humana), ('sessoes sem mensagem do humano (automaticas/legado)', lambda k: k not in sess_humana)):
        ks = [k for k in sess_tok if cond(k)]
        tt_ = sum(v['entrada'] + v['lido'] + v['grav'] + v['saida'] for k, v in sess_tok.items() if cond(k))
        ch_ = sum(v['chamadas'] for k, v in sess_tok.items() if cond(k))
        out('  %-55s sessoes=%3d chamadas=%6d tokens=%14d (%.1f%% do total)' % (nome, len(ks), ch_, tt_, 100.0 * tt_ / max(1, sum(soma_s))))
    out('  por dia (chamadas / tokens totais):')
    for d in dias:
        cs = [c for c in ch_sem if dia_local(c['ts']) == d]
        out('    %s %6d chamadas  %14d tokens' % (d, len(cs), sum(c['ent'] + c['lido'] + c['grav'] + c['saida'] for c in cs)))

    out('\n[2] CONTEXTO POR CHAMADA (entrada + cache lido + cache gravado)')
    out('  principais: n=%d mediana=%d p90=%d max=%d' % (len(ctx_pr), med(ctx_pr), pct(ctx_pr, .9), max(ctx_pr) if ctx_pr else 0))
    out('  subagentes: n=%d mediana=%d p90=%d max=%d' % (len(ctx_sub), med(ctx_sub), pct(ctx_sub, .9), max(ctx_sub) if ctx_sub else 0))

    # com --mapa, as pastas do mesmo projeto somam as sessoes (a mediana e sobre todas elas); sem ele, por pasta
    ctx0_por_proj = collections.defaultdict(list)
    for p, v in ctx0_proj.items():
        ctx0_por_proj[projeto_de(p, mapa)].extend(v)
    todos0 = [x for v in ctx0_proj.values() for x in v]
    out('\n[3] ctx0 POR %s (1a chamada da sessao principal; so sessoes do humano cuja 1a chamada cai na semana)' % (
        'PASTA (sem --mapa)' if mapa is None else 'PROJETO (pelo mapa)'))
    for p, v in sorted(ctx0_por_proj.items()):
        out('  %-60s sessoes=%2d  ctx0 mediana=%7d min=%7d max=%7d' % (p, len(v), med(v), min(v), max(v)))
    out('  sessoes do humano: n=%d mediana=%d | todas as sessoes principais (sem filtro): n=%d mediana=%d' % (len(todos0), med(todos0), len(ctx0_todas), med(ctx0_todas)))
    sem_mapa = sorted({c['projeto'] for c in ch_sem if mapa is not None and c['projeto'].lower() not in mapa})
    if mapa is not None:
        out('  pastas com chamada na semana e fora do mapa (viram "desconhecido"): %d %s' % (len(sem_mapa), sem_mapa))

    out('\n[4] MENSAGENS DO HUMANO POR SESSAO (sessoes com algum evento humano na semana: %d)' % len(hum_sess))
    tot_c = collections.Counter()
    for n in hum_sess.values():
        tot_c.update(n)
    out('  totais por classe:', dict(sorted(tot_c.items())), ' (H digitada, HQ digitada com Claude ocupado, HC comando /, HA resposta/recusa, HI interrupcao)')
    out('  estrito (H+HQ+HC) por sessao: n=%d mediana=%.1f p90=%d max=%d  soma=%d' % (
        len(msgs_estrito), med(msgs_estrito), pct(msgs_estrito, .9), max(msgs_estrito) if msgs_estrito else 0, sum(msgs_estrito)))
    out('  distribuicao (n de mensagens estritas -> n de sessoes):', dict(sorted(collections.Counter(min(x, 10) for x in msgs_estrito).items())), '(10 = 10 ou mais)')
    out('  sessoes iniciadas pelo humano na semana (1o evento humano estrito dentro da semana): %d' % inicios)

    out('\n[5] RETORNOS (evento humano depois de > 30 min sem evento na sessao; o 1o evento humano da sessao nao conta)')
    for k in sorted(retornos):
        out('  conjunto %-8s %-26s: %d' % (k[0], k[1], retornos[k]))
    out('  por dia (v2, amplo):', {d: retornos_dia.get(('amplo', d), 0) for d in dias})
    out('  por dia (v2, estrito):', {d: retornos_dia.get(('estrito', d), 0) for d in dias})
    out('  variantes de limiar (v2): ' + ', '.join('%s %dmin=%d' % (c, lim, retornos_var[(c, lim)]) for c in ('estrito', 'amplo') for lim in LIM_RETORNO_VAR))
    faixas = [(0, 60), (60, 300), (300, 1200), (1200, 1800), (1800, 3600), (3600, 10800), (10800, 10 ** 9)]
    nomes = ['<1min', '1-5min', '5-20min', '20-30min', '30-60min', '1-3h', '>3h']
    out('  lacuna antes de cada evento humano (v2, amplo, exceto o 1o): ' + ', '.join('%s=%d' % (n, sum(1 for g in gaps_h if lo <= g < hi)) for n, (lo, hi) in zip(nomes, faixas)))

    def resumo_trechos(tr, nome):
        tr = sorted(tr, reverse=True)
        durs = [x[0] for x in tr]
        out('  %-34s n=%d  mediana=%s  p90=%s  MAIOR=%s' % (nome, len(tr), hm(med(durs)), hm(pct(durs, .9)), hm(tr[0][0]) if tr else '-'))
        out('      >=10min: %d  >=30min: %d  >=1h: %d  >=2h: %d  >=4h: %d  >=8h: %d' % tuple(sum(1 for x in durs if x >= lim) for lim in (600, 1800, 3600, 7200, 14400, 28800)))
        out('      5 maiores (duracao, dia local, hora local, chamadas/eventos de trabalho no trecho):')
        for dur, h, n in tr[:5]:
            out('        %s  %s %s  %s' % (hm(dur), dia_local(h), dt.datetime.fromtimestamp(h).strftime('%H:%M'), n if n else '-'))
    out('\n[6] MAIOR TRECHO TRABALHANDO SEM O HUMANO')
    resumo_trechos(trechos_A, 'A literal (fronteira: humano estrito)')
    resumo_trechos(trechos_B, 'B (fronteira: presenca ampla)')
    resumo_trechos(trechos_C, 'C continuo (lacuna <= %g min)' % (GAP_CONTINUO / 60))
    resumo_trechos(trechos_G, 'G global (todas as sessoes, <= %g min)' % (GAP_CONTINUO / 60))

    out('\n[7] CONCORRENCIA (sessoes distintas com evento de trabalho do Claude, principal+subagentes)')
    out('  sessoes do humano: %d de %d sessoes com trabalho na semana; com >= 10 min de trabalho: %d' % (
        len([k for k in sess_humana if dur_trab.get(k, 0) > 0]), len([k for k in por_sessao if dur_trab.get(k, 0) > 0]),
        len([k for k in sess_humana if dur_trab.get(k, 0) >= 600])))
    out('  dia        | max sessoes na janela de 10 min: bruta / do humano / do humano>=10min | horas c/ 2+ (do humano): solto / janela 10min | pico por span (bruta, como o session-report)')
    for d in dias:
        out('  %s | %26d / %9d / %14d | %24d / %12d | %d' % (d, c_bruta[0][d], c_hum[0][d], c_hum10[0][d], c_hum[1].get(d, 0), c_hum[2].get(d, 0), pico_span[d]))
    out('  semana: max bruta=%d  max do humano=%d  max do humano>=10min=%d | horas 2+ do humano: solto=%d janela=%d | horas 2+ bruta: solto=%d janela=%d' % (
        max(c_bruta[0].values()), max(c_hum[0].values()), max(c_hum10[0].values()), sum(c_hum[1].values()), sum(c_hum[2].values()), sum(c_bruta[1].values()), sum(c_bruta[2].values())))
    out('  horas de relogio com algum trabalho na semana: %d (bruta) / %d (do humano)' % (c_bruta[3], c_hum[3]))

    out('\n[8] SENTADAS (mensagens do humano de todas as sessoes juntas; intervalo <= 20 min)')
    for nome, bl in (('estrito (H+HQ+HC)', sent_e), ('amplo (+ respostas/recusas/interrupcoes)', sent_a)):
        out('  %s: blocos=%d  duracao somada=%s  blocos de 1 mensagem=%d  maior bloco=%s' % (
            nome, len(bl), hm(sum(b[1] - b[0] for b in bl)), sum(1 for b in bl if b[2] == 1), hm(max((b[1] - b[0] for b in bl), default=0))))
        por_d = collections.defaultdict(lambda: [0, 0.0, 0])
        for b in bl:
            x = por_d[dia_local(b[0])]
            x[0] += 1
            x[1] += b[1] - b[0]
            x[2] += b[2]
        out('    por dia (blocos / duracao / mensagens):', {d: (por_d[d][0], hm(por_d[d][1]), por_d[d][2]) for d in dias})
    for m in LIM_SENTADA_VAR:
        bl = sent_var[m]
        out('  variante estrito, intervalo <= %d min: blocos=%d  duracao somada=%s' % (m, len(bl), hm(sum(b[1] - b[0] for b in bl))))

    out('\n[9] CHAMADAS POR MODELO (message.id unicos na semana; principal | subagente)')
    todos = sorted(set(por_modelo) | set(por_modelo_sub))
    for m in todos:
        out('  %-28s principal %6d | subagente %6d' % (m, por_modelo[m], por_modelo_sub[m]))
    out('  por modelo e esforco (principal+subagente):', dict(sorted(por_modelo_esf.items(), key=lambda x: -x[1])))
    out('  por esforco (principal+subagente):', dict(sorted(esf_tot.items(), key=lambda x: -x[1])))
    fora_colunas = tok['total']['chamadas'] - sum(por_modelo[m] + por_modelo_sub[m] for m, _ in MODELOS)
    out('  chamadas de modelos sem coluna no CSV: %d' % fora_colunas)
    out('  mensagens sinteticas de erro da API por (dia, quota, arquivo):', {k: v for k, v in sorted(ctrl['erro_api'].items()) if dias[0] <= k[0] <= dias[-1]})
    out('  limites batidos (erro sintetico da API com quotaLimits) na semana: %d (arquivos principais %d, subagentes %d)' % (
        len(lim_sem), sum(1 for x in lim_sem if not x), sum(1 for x in lim_sem if x)))

    out('\n[10] TEMPO DO SCRIPT: total %.1fs (leitura dos arquivos %.1fs)' % (dur_prog, t_leitura))

    # linha semanal (CSV): so inteiros, a data ISO e os rotulos pc e versao; '' onde nao ha dado
    tot = tok['total']

    def mediana_int(v):                   # mediana inteira (trunca, como no prototipo)
        return int(med(v)) if v else ''

    def maior_min(tr):                    # maior trecho em minutos arredondados
        return round(max(x[0] for x in tr) / 60) if tr else ''

    linha = collections.OrderedDict()
    linha['semana_inicio'] = dias[0]
    linha['sessoes_com_chamada'] = len(sess_tok)
    linha['chamadas'] = tot['chamadas']
    linha['chamadas_subagente'] = tok['sub']['chamadas']
    linha['tok_entrada'] = tot['entrada']
    linha['tok_cache_lido'] = tot['cache_lido']
    linha['tok_cache_gravado'] = tot['cache_gravado']
    linha['tok_saida'] = tot['saida']
    linha['ctx_mediano_principal'] = mediana_int(ctx_pr)
    linha['ctx0_mediano'] = mediana_int(todos0)
    linha['msgs_humano_estrito'] = sum(msgs_estrito)
    linha['respostas_perguntas'] = tot_c['HA']
    linha['retornos_30min_estrito'] = retornos[('estrito', 'v2_so_trabalho_e_humano')]
    linha['retornos_30min_amplo'] = retornos[('amplo', 'v2_so_trabalho_e_humano')]
    linha['maior_trecho_A_min'] = maior_min(trechos_A)
    linha['maior_trecho_B_min'] = maior_min(trechos_B)
    linha['maior_trecho_C_min'] = maior_min(trechos_C)
    linha['conc_max_humano'] = max(conc_dia.values())
    linha['conc_max_bruta'] = max(c_bruta[0].values())
    linha['horas_2mais_humano_janela10'] = sum(h2_estrito.values())
    linha['sentadas_estrito'] = len(sent_e)
    linha['sentadas_min_estrito'] = round(sum(b[1] - b[0] for b in sent_e) / 60)
    for m, col in MODELOS:
        linha[col] = por_modelo[m] + por_modelo_sub[m]
    linha['maior_trecho_global_min'] = maior_min(trechos_G)
    linha['trecho_global_mediana_min'] = round(med([x[0] for x in trechos_G]) / 60) if trechos_G else ''
    for m in LIM_RETORNO_VAR:
        linha['retornos_%dmin_estrito' % m] = retornos_var[('estrito', m)]
    for m in LIM_SENTADA_VAR:
        linha['sentadas_%dmin_blocos' % m] = len(sent_var[m])
        linha['sentadas_%dmin_min' % m] = round(sum(b[1] - b[0] for b in sent_var[m]) / 60)
    linha['chamadas_esforco_max'] = esf_tot['max']
    linha['chamadas_esforco_xhigh'] = esf_tot['xhigh']
    linha['limites_batidos'] = len(lim_sem)
    linha['pct_limite_semanal'] = pct_semanal(a.limites, t_ini, t_fim)   # so a statusline traz o numero
    linha['pc'] = a.pc
    linha['versao'] = VERSAO
    assert list(linha) == CABECALHO, 'colunas fora do esquema'

    print('\nLINHA SEMANAL:')
    print(','.join(linha.keys()))
    print(','.join(str(v) for v in linha.values()))
    avisos = []
    if (a.retorno_min, a.sentada_min, a.gap_continuo_min) != (30, 20, 15):
        avisos.append('limiares diferentes do padrao (retorno 30, sentada 20, trecho 15 min): os nomes das colunas '
                      '(retornos_30min_*, sentadas_estrito, maior_trecho_*) seguem os do padrao')
    if a.dias != 7:
        avisos.append('--dias diferente de 7: a linha nao e de uma semana inteira')
    if d0.weekday() != 1:
        avisos.append('--inicio nao e terca-feira: a semana do CSV vai de terca a segunda')
    if fora_colunas:
        avisos.append('%d chamadas de modelos sem coluna no CSV (a soma das colunas de modelo fica menor que "chamadas")' % fora_colunas)
    for av in avisos:
        print('AVISO: ' + av, file=sys.stderr)
    erro_csv = []
    if a.csv:
        i_pc = CABECALHO.index('pc')
        if not gravar_csv(a.csv, CABECALHO, [[str(v) for v in linha.values()]],
                          lambda r: len(r) > i_pc and r[0] == dias[0] and r[i_pc] == a.pc, (0, i_pc), ('pct_limite_semanal',)):
            erro_csv.append(a.csv)
    if a.ctx0_csv:
        novas = [[dias[0], p, str(len(v)), str(int(med(v))), str(min(v)), str(max(v))] for p, v in sorted(ctx0_por_proj.items())]
        if not gravar_csv(a.ctx0_csv, CABECALHO_CTX0, novas, lambda r: r[0] == dias[0], (0, 1)):
            erro_csv.append(a.ctx0_csv)
    for f in erro_csv:
        print('ERRO: o cabecalho do CSV existente difere do esperado; nao gravei: ' + f, file=sys.stderr)
    return 3 if erro_csv else 0


if __name__ == '__main__':
    sys.exit(main())
